"""SpecDocService converts Word and presentation members through an injected pool."""

from __future__ import annotations

import io
import zipfile
from unittest.mock import MagicMock

import pytest
from docx import Document

from doc3gpp.models.spec import SpecVersion
from doc3gpp.parsers.docx_converter import (
    InlineDocxPool,
    ParagraphBlock,
    PythonDocxNotInstalledError,
)
from doc3gpp.services.docx_convert_pool import DocxConvertPool
from doc3gpp.services.spec_doc_service import SpecDocService
from tests.fixtures.presentation_builders import build_presentation_with_content


def _docx_bytes(heading: str, body: str) -> bytes:
    doc = Document()
    doc.add_heading(heading, level=1)
    doc.add_paragraph(body)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _multi_docx_zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("38331-j00.docx", _docx_bytes("1 Scope", "scope body"))
        archive.writestr("38331-j10.docx", _docx_bytes("2 References", "reference body"))
        archive.writestr("38331-j20.docx", _docx_bytes("3 Definitions", "definitions body"))
        archive.writestr("slides.pptx", build_presentation_with_content())
    return buf.getvalue()


class _RecordingPool:
    """Records each call; optionally replaces the second member's outcome with ``error``."""

    def __init__(self, error: BaseException | None = None) -> None:
        self._error = error
        self.calls: list[tuple[str, list[str]]] = []

    def convert_members(self, kind, members):
        self.calls.append((kind, [name for name, _ in members]))
        outcomes = InlineDocxPool().convert_members(kind, members)
        if self._error is not None:
            outcomes[1] = self._error
        return outcomes


def _service(tmp_path, **kwargs) -> SpecDocService:
    spec_repo = MagicMock()
    spec_repo.list_versions.return_value = [
        SpecVersion("38.331", "18.5.0", "https://example.test/doc.zip", "Rel-18")
    ]
    repo = MagicMock()
    repo.get_source.return_value = None
    raw = _multi_docx_zip()
    return SpecDocService(
        spec_repo=spec_repo,
        doc_repo=repo,
        inventory_repo=MagicMock(),
        fetcher=lambda _: raw,
        cache_dir=tmp_path,
        **kwargs,
    )


def test_pool_and_default_produce_equal_prepared_output(tmp_path) -> None:
    inline = _service(tmp_path / "inline").prepare_parse("38.331")
    with DocxConvertPool(2) as pool:
        pooled = _service(tmp_path / "pool", docx_pool=pool).prepare_parse("38.331")

    assert pooled.toc == inline.toc
    assert pooled.drafts == inline.drafts
    assert pooled.ordered_files == inline.ordered_files
    assert pooled.presentations == inline.presentations


def test_pool_receives_blocks_kind_and_members_in_zip_order(tmp_path) -> None:
    pool = _RecordingPool()
    _service(tmp_path, docx_pool=pool).prepare_parse("38.331")

    assert pool.calls == [
        (
            "blocks",
            [
                "38331-j00.docx",
                "38331-j10.docx",
                "38331-j20.docx",
                "slides.pptx",
            ],
        )
    ]


def test_presentation_preparation_retains_full_markdown_without_toc_metadata(tmp_path) -> None:
    prepared = _service(tmp_path).prepare_parse("38.331")

    presentation = prepared.presentations[0]
    slide_file = next(
        file for file in prepared.ordered_files if file.source_file == "slides.pptx"
    )
    slide_chunks = [
        draft for draft in prepared.drafts if draft.source_file == "slides.pptx"
    ]

    assert presentation.source_file == "slides.pptx"
    assert presentation.file_order == slide_file.file_order
    assert "speaker secret" in presentation.markdown
    assert len(slide_file.blocks) == 1
    assert isinstance(slide_file.blocks[0], ParagraphBlock)
    assert slide_chunks
    assert all(chunk.sections is None and chunk.tables is None for chunk in slide_chunks)
    assert all(entry.source_file != "slides.pptx" for entry in prepared.toc.entries)
    assert all(file.source_file != "slides.pptx" for file in prepared.toc.files)


def test_pool_exception_outcome_is_raised_unchanged(tmp_path) -> None:
    error = RuntimeError("boom")
    service = _service(tmp_path, docx_pool=_RecordingPool(error))

    prepared = service.prepare_parse("38.331")
    assert prepared.error is error

    with pytest.raises(RuntimeError, match="boom") as exc_info:
        service.persist_prepared_parse(prepared)

    assert exc_info.value is error


def test_pool_import_error_outcome_is_raised_unchanged(tmp_path) -> None:
    error = PythonDocxNotInstalledError()
    service = _service(tmp_path, docx_pool=_RecordingPool(error))

    prepared = service.prepare_parse("38.331")
    assert prepared.error is error

    with pytest.raises(PythonDocxNotInstalledError) as exc_info:
        service.persist_prepared_parse(prepared)

    assert exc_info.value is error

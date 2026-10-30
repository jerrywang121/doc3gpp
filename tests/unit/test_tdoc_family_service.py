"""Unit tests for WID, LS, and discussion family extraction."""

from __future__ import annotations

import hashlib
import io
import logging
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest

from doc3gpp.models.tdoc import TDoc
from doc3gpp.models.tdoc_content_families import TDOC_CONTENT_STORES
from doc3gpp.parsers.tdoc_family_main import TDocMainDocumentError
from doc3gpp.parsers.tdoc_zip import TDocZipLimitError, read_markdown_archive
from doc3gpp.scraping.cache_keys import derive_cache_file
from doc3gpp.scraping.tdoc_zip_source import DownloadedZip
from doc3gpp.services import tdoc_family_service
from doc3gpp.services.tdoc_errors import TDocNotYetOnFTPError
from doc3gpp.services.tdoc_family_service import (
    FamilyExtractResult,
    PreparedFamilyExtract,
    TDocFamilyService,
)
from doc3gpp.storage.db.migrate import create_schema
from doc3gpp.storage.db.models import TDocORM
from doc3gpp.storage.db.session import get_session_factory
from doc3gpp.storage.repositories.tdoc_content_sql import (
    SQLAlchemyTDocContentRepository,
)


def _docx_available() -> bool:
    try:
        from docx import Document  # noqa: F401
    except ImportError:
        return False
    return True


pytestmark = pytest.mark.skipif(
    not _docx_available(),
    reason="python-docx not installed; install with `pip install doc3gpp[extract]`",
)

STORES = {store.key: store for store in TDOC_CONTENT_STORES}
WID = STORES["wid"]
LS = STORES["ls"]
TDOC_ID = "R5-123456"
FTP_URL = "tsg_ran/WG5_Test_Ex-T1/TSGR5_100/Docs/R5-123456.zip"
CANONICAL = "tsg_ran/wg5_test_ex-t1/tsgr5_100/docs/r5-123456.zip"
MAIN = "R5-123456.docx"


def _save_docx(document) -> bytes:
    payload = io.BytesIO()
    document.save(payload)
    return payload.getvalue()


def _make_docx(text: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    return _save_docx(document)


def _make_empty_docx() -> bytes:
    from docx import Document

    return _save_docx(Document())


def _make_zip(files: list[tuple[str, bytes]]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, payload in files:
            archive.writestr(name, payload)
    return output.getvalue()


def _tdoc() -> TDoc:
    return TDoc(tdoc_id=TDOC_ID, ftp_url=FTP_URL, type="WID revised")


class _FakeCache:
    def __init__(self) -> None:
        self.puts: list[tuple[str, bytes, str]] = []

    def put_bytes(self, key: str, payload: bytes, subdir: str) -> Path:
        self.puts.append((key, payload, subdir))
        return Path(key)


class _Downloads:
    """Serves the configured ZIP bytes through ``download_tdoc_zip`` and counts calls."""

    def __init__(self, directory: Path) -> None:
        self._directory = directory
        self.zip_bytes: bytes = b""
        self.calls = 0

    def serve(self, zip_bytes: bytes) -> None:
        self.zip_bytes = zip_bytes

    def download(
        self,
        tdoc,
        scraper,
        cache,
        primary_url=None,
        *,
        ftp_url=None,
        max_bytes=0,
        require_provenance=False,
    ):
        del tdoc, scraper, cache, ftp_url, max_bytes, require_provenance
        self.calls += 1
        path = self._directory / "download.zip"
        path.write_bytes(self.zip_bytes)
        return DownloadedZip(
            path=path,
            url=primary_url,
            payload=self.zip_bytes,
        )


@pytest.fixture()
def downloads(monkeypatch, tmp_path) -> _Downloads:
    served = _Downloads(tmp_path)
    monkeypatch.setattr(tdoc_family_service, "download_tdoc_zip", served.download)
    return served


@pytest.fixture()
def family_db(sqlite_env) -> None:
    create_schema("main")
    with get_session_factory()() as session:
        session.add(TDocORM(tdoc_id=TDOC_ID))
        session.commit()


def _repo() -> SQLAlchemyTDocContentRepository:
    return SQLAlchemyTDocContentRepository()


def _service(cache: _FakeCache, *, max_documents: int = 100) -> TDocFamilyService:
    return TDocFamilyService(
        scraper=object(),
        cache=cache,
        content_repo=_repo(),
        max_tdoc_size_bytes=0,
        max_documents_per_zip=max_documents,
        max_uncompressed_size_bytes=0,
    )


def _run(service: TDocFamilyService, store, *, force: bool = False) -> FamilyExtractResult:
    prepared = service.prepare_extract(_tdoc(), store, force=force)
    if isinstance(prepared, FamilyExtractResult):
        return prepared
    return service.persist_prepared_extract(prepared)


def test_stores_main_and_annexes_in_order(family_db, downloads) -> None:
    downloads.serve(
        _make_zip(
            [
                ("annex/b.docx", _make_docx("annex b body")),
                (MAIN, _make_docx("main body")),
                ("annex/a.docx", _make_docx("annex a body")),
            ]
        )
    )

    result = _run(_service(_FakeCache()), WID)

    assert result.extract_meta is not None
    assert result.extract_meta.document_count == 3
    assert result.extract_meta.doc_filename == MAIN
    assert [document.source_file for document in result.documents] == [
        MAIN,
        "annex/a.docx",
        "annex/b.docx",
    ]
    assert [document.markdown_text for document in result.documents] == [
        "main body\n",
        "annex a body\n",
        "annex b body\n",
    ]
    assert result.skipped_sources == ()
    assert result.complete is True
    assert result.from_cache is False
    assert _repo().list_documents(WID, CANONICAL) == result.documents


def test_failing_annex_is_skipped_recorded_and_not_fatal(family_db, downloads, caplog) -> None:
    downloads.serve(
        _make_zip(
            [
                (MAIN, _make_docx("main body")),
                ("annex/bad.docx", b"not a docx"),
                ("annex/a.docx", _make_docx("annex a body")),
            ]
        )
    )

    with caplog.at_level(logging.WARNING):
        result = _run(_service(_FakeCache()), WID)

    assert result.skipped_sources == ("annex/bad.docx",)
    assert result.extract_meta is not None
    assert result.extract_meta.skipped_sources == ("annex/bad.docx",)
    assert result.extract_meta.document_count == 2
    assert result.complete is False
    assert [document.source_file for document in result.documents] == [
        MAIN,
        "annex/a.docx",
    ]
    assert [document.source_file for document in _repo().list_documents(WID, CANONICAL)] == [
        MAIN,
        "annex/a.docx",
    ]
    assert any(
        record.levelno == logging.WARNING and "annex/bad.docx" in record.getMessage()
        for record in caplog.records
    )


def test_failing_main_fails_and_keeps_previous_content(family_db, downloads) -> None:
    from doc3gpp.services.tdoc_zip_inventory import PreparedInventoryFailure

    service = _service(_FakeCache())
    downloads.serve(
        _make_zip(
            [
                (MAIN, _make_docx("main body")),
                ("annex/a.docx", _make_docx("annex a body")),
            ]
        )
    )
    _run(service, WID)

    downloads.serve(
        _make_zip(
            [
                (MAIN, b"not a docx"),
                ("annex/a.docx", _make_docx("new annex body")),
            ]
        )
    )
    prepared = service.prepare_extract(_tdoc(), WID, force=True)
    assert isinstance(prepared, PreparedInventoryFailure)
    assert isinstance(prepared.error, zipfile.BadZipFile)

    listed = _repo().list_documents(WID, CANONICAL)
    assert [document.source_file for document in listed] == [MAIN, "annex/a.docx"]
    assert [document.markdown_text for document in listed] == [
        "main body\n",
        "annex a body\n",
    ]
    meta = _repo().get_extract_meta(CANONICAL)
    assert meta is not None
    assert meta.document_count == 2


def test_unsupported_main_fails_with_word_format_reason(family_db, downloads) -> None:
    from doc3gpp.services.tdoc_zip_inventory import PreparedInventoryFailure

    downloads.serve(_make_zip([(f"{TDOC_ID}.pdf", b"unsupported document")]))

    prepared = _service(_FakeCache()).prepare_extract(_tdoc(), WID)

    assert isinstance(prepared, PreparedInventoryFailure)
    assert isinstance(prepared.error, TDocMainDocumentError)
    assert "is not .docx" in str(prepared.error)


def test_family_main_in_folder_stored_with_full_path(family_db, downloads) -> None:
    main_in_folder = f"{TDOC_ID}/{MAIN}"
    downloads.serve(
        _make_zip(
            [
                ("annex/a.docx", _make_docx("annex a body")),
                (main_in_folder, _make_docx("main body")),
            ]
        )
    )

    result = _run(_service(_FakeCache()), WID)

    assert result.extract_meta is not None
    assert result.extract_meta.doc_filename == main_in_folder
    assert result.documents[0].source_file == main_in_folder
    assert [document.source_file for document in result.documents] == [
        main_in_folder,
        "annex/a.docx",
    ]


def test_zero_text_main_is_stored_and_complete(family_db, downloads) -> None:
    downloads.serve(_make_zip([(MAIN, _make_empty_docx())]))

    result = _run(_service(_FakeCache()), WID)

    assert result.extract_meta is not None
    assert result.extract_meta.document_count == 1
    assert [document.markdown_text for document in result.documents] == [""]
    assert [document.content_hash for document in result.documents] == [
        hashlib.sha256(b"").hexdigest()
    ]
    assert result.complete is True
    assert _repo().list_documents(WID, CANONICAL) == result.documents


def test_cache_hit_skips_download_when_complete(family_db, downloads, monkeypatch) -> None:
    cache = _FakeCache()
    service = _service(cache)
    downloads.serve(
        _make_zip(
            [
                (MAIN, _make_docx("main body")),
                ("annex/a.docx", _make_docx("annex a body")),
            ]
        )
    )
    persisted = _run(service, WID)

    def must_not_download(*args, **kwargs):
        raise AssertionError("complete SQL extraction must not download the ZIP")

    monkeypatch.setattr(tdoc_family_service, "download_tdoc_zip", must_not_download)
    cached = service.prepare_extract(_tdoc(), WID)

    assert isinstance(cached, FamilyExtractResult)
    assert cached.from_cache is True
    assert cached.complete is True
    assert cached.documents == persisted.documents
    assert replace(cached.extract_meta, extracted_at=None) == persisted.extract_meta
    assert len(cache.puts) == 1


def test_force_redownloads_even_when_complete(family_db, downloads) -> None:
    service = _service(_FakeCache())
    downloads.serve(_make_zip([(MAIN, _make_docx("main body"))]))
    _run(service, WID)
    assert downloads.calls == 1

    prepared = service.prepare_extract(_tdoc(), WID, force=True)

    assert isinstance(prepared, PreparedFamilyExtract)
    assert downloads.calls == 2


def test_incomplete_extraction_is_retried_by_next_normal_parse(family_db, downloads) -> None:
    service = _service(_FakeCache())
    downloads.serve(
        _make_zip(
            [
                (MAIN, _make_docx("main body")),
                ("annex/bad.docx", b"not a docx"),
            ]
        )
    )
    first = _run(service, WID)
    assert first.skipped_sources == ("annex/bad.docx",)
    assert first.complete is False

    downloads.calls = 0
    prepared = service.prepare_extract(_tdoc(), WID)

    assert isinstance(prepared, PreparedFamilyExtract)
    assert downloads.calls == 1


def test_archive_contains_every_converted_member(family_db, downloads) -> None:
    cache = _FakeCache()
    downloads.serve(
        _make_zip(
            [
                ("annex/b.docx", _make_docx("annex b body")),
                (MAIN, _make_docx("main body")),
                ("annex/a.docx", _make_docx("annex a body")),
            ]
        )
    )

    prepared = _service(cache).prepare_extract(_tdoc(), WID)

    assert isinstance(prepared, PreparedFamilyExtract)
    assert len(cache.puts) == 1
    key, payload, subdir = cache.puts[0]
    assert key == derive_cache_file(FTP_URL)
    assert subdir == "markdown"
    archive = read_markdown_archive(payload)
    assert archive.ftp_url == CANONICAL
    assert archive.main_source_file == MAIN
    assert [document.source_file for document in archive.documents] == [
        MAIN,
        "annex/a.docx",
        "annex/b.docx",
    ]
    assert [document.markdown_text for document in archive.documents] == [
        "main body\n",
        "annex a body\n",
        "annex b body\n",
    ]


def test_prepare_writes_no_sql_until_persist(family_db, downloads) -> None:
    service = _service(_FakeCache())
    downloads.serve(
        _make_zip(
            [
                (MAIN, _make_docx("main body")),
                ("annex/a.docx", _make_docx("annex a body")),
            ]
        )
    )

    prepared = service.prepare_extract(_tdoc(), WID)

    assert isinstance(prepared, PreparedFamilyExtract)
    assert _repo().list_documents(WID, CANONICAL) == ()
    assert _repo().get_extract_meta(CANONICAL) is None

    result = service.persist_prepared_extract(prepared)

    assert _repo().list_documents(WID, CANONICAL) == result.documents
    assert [document.source_file for document in result.documents] == [
        MAIN,
        "annex/a.docx",
    ]
    meta = _repo().get_extract_meta(CANONICAL)
    assert meta is not None
    assert meta.doc_filename == MAIN
    assert meta.document_count == 2


def test_zip_document_limit_raises_limit_error(family_db, downloads) -> None:
    from doc3gpp.services.tdoc_zip_inventory import PreparedInventoryFailure

    downloads.serve(
        _make_zip(
            [
                (MAIN, _make_docx("main body")),
                ("annex/a.docx", _make_docx("annex a body")),
            ]
        )
    )

    prepared = _service(_FakeCache(), max_documents=1).prepare_extract(_tdoc(), WID)

    assert isinstance(prepared, PreparedInventoryFailure)
    assert isinstance(prepared.error, TDocZipLimitError)
    assert "document-count limit exceeded" in str(prepared.error)


def test_cr_store_is_rejected_before_download(family_db, monkeypatch) -> None:
    def must_not_download(*args, **kwargs):
        raise AssertionError("a non-family store must not download")

    monkeypatch.setattr(tdoc_family_service, "download_tdoc_zip", must_not_download)

    with pytest.raises(ValueError, match="not a family store"):
        _service(_FakeCache()).prepare_extract(_tdoc(), STORES["cr"])


def test_tdoc_without_ftp_url_is_rejected_before_download(family_db, monkeypatch) -> None:
    def must_not_download(*args, **kwargs):
        raise AssertionError("a TDoc without ftp_url must not download")

    monkeypatch.setattr(tdoc_family_service, "download_tdoc_zip", must_not_download)

    with pytest.raises(TDocNotYetOnFTPError) as excinfo:
        _service(_FakeCache()).prepare_extract(TDoc(tdoc_id=TDOC_ID), WID)

    assert type(excinfo.value) is TDocNotYetOnFTPError
    assert excinfo.value.tdoc_id == TDOC_ID


def test_ls_store_persists_into_ls_table(family_db, downloads) -> None:
    downloads.serve(_make_zip([(MAIN, _make_docx("main body"))]))

    result = _run(_service(_FakeCache()), LS)

    assert result.store is LS
    assert _repo().list_documents(LS, CANONICAL) == result.documents
    assert _repo().list_documents(WID, CANONICAL) == ()



def test_prepare_attaches_outer_zip_inventory_for_non_docx_entries(
    family_db, downloads
) -> None:
    downloads.serve(
        _make_zip(
            [
                (MAIN, _make_docx("main body")),
                ("support/reference.pdf", b"not converted"),
            ]
        )
    )

    prepared = _service(_FakeCache()).prepare_extract(_tdoc(), WID, force=True)

    assert prepared.inventory is not None
    assert prepared.inventory.zip_files == (MAIN, "support/reference.pdf")

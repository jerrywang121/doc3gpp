"""DOCX conversion pool changes no parse output.

Each path runs once in-process and once through a real ``DocxConvertPool(2)``
with spawned worker processes. The results must compare equal. The recording
wrapper proves the pool was the converter and that worker processes were alive
during the call.
"""

from __future__ import annotations

import io
import os
import zipfile
from unittest.mock import MagicMock

from docx import Document

from doc3gpp.models.spec import SpecVersion
from doc3gpp.parsers.tdoc_zip_convert import convert_tdoc_zip
from doc3gpp.services.docx_convert_pool import DocxConvertPool
from doc3gpp.services.spec_doc_service import SpecDocService
from doc3gpp.settings.loader import get_settings
from tests.fixtures.docx_builders import build_golden_branches_docx, build_golden_docx

SPEC_ID = "38.331"
TDOC_ID = "R5-123456"
IDENTITY = "tsg_ran/wg5_test_ex-t1/tsgr5_100/docs/r5-123456.zip"


def _heading_docx(heading: str, body: str) -> bytes:
    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(body)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _golden_with_extra_paragraphs() -> bytes:
    document = Document(io.BytesIO(build_golden_docx()))
    for index in range(5):
        document.add_paragraph(f"Extra paragraph {index}.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _documents() -> dict[str, bytes]:
    return {
        "references": _heading_docx("2 References", "reference body"),
        "golden": build_golden_docx(),
        "branches": build_golden_branches_docx(),
        "golden_extra": _golden_with_extra_paragraphs(),
    }


def _zip(files: list[tuple[str, bytes]]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, payload in files:
            archive.writestr(name, payload)
    return buffer.getvalue()


class _RecordingPool:
    """Delegates to ``inner`` and records each call and the pool's live worker PIDs."""

    def __init__(self, inner) -> None:
        self._inner = inner
        self.calls: list[tuple[str, list[str]]] = []
        self.worker_pids: set[int] = set()

    def convert_members(self, kind, members):
        self.calls.append((kind, [name for name, _ in members]))
        outcomes = self._inner.convert_members(kind, members)
        self.worker_pids.update(self._inner._executor._processes)
        return outcomes


def _pid_alive(pid: int) -> bool:
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    return True


def _spec_service(tmp_path, docx_pool=None) -> SpecDocService:
    spec_repo = MagicMock()
    spec_repo.list_versions.return_value = [
        SpecVersion(SPEC_ID, "18.5.0", "https://example.test/38331-h50.zip", "Rel-18")
    ]
    doc_repo = MagicMock()
    doc_repo.get_source.return_value = None
    documents = _documents()
    raw = _zip(
        [
            ("38331-j00.docx", documents["references"]),
            ("38331-j10.docx", documents["golden"]),
            ("38331-j20.docx", documents["branches"]),
            ("38331-j30.docx", documents["golden_extra"]),
        ]
    )
    return SpecDocService(
        spec_repo=spec_repo,
        doc_repo=doc_repo,
        settings=get_settings(),
        fetcher=lambda _: raw,
        cache_dir=tmp_path,
        docx_pool=docx_pool,
    )


def test_spec_doc_prepare_parse_pool_matches_in_process(tmp_path) -> None:
    inline = _spec_service(tmp_path / "inline").prepare_parse(SPEC_ID)

    with DocxConvertPool(2) as pool:
        recording = _RecordingPool(pool)
        pooled = _spec_service(tmp_path / "pooled", docx_pool=recording).prepare_parse(SPEC_ID)

    assert recording.calls == [
        ("blocks", ["38331-j00.docx", "38331-j10.docx", "38331-j20.docx", "38331-j30.docx"])
    ]
    assert recording.worker_pids
    assert inline.toc.entries
    assert inline.drafts
    assert [f.source_file for f in inline.ordered_files] != [
        "38331-j00.docx",
        "38331-j10.docx",
        "38331-j20.docx",
        "38331-j30.docx",
    ]
    assert pooled.toc == inline.toc
    assert pooled.drafts == inline.drafts
    assert pooled.ordered_files == inline.ordered_files


def test_convert_tdoc_zip_pool_matches_in_process(tmp_path) -> None:
    documents = _documents()
    zip_bytes = _zip(
        [
            (f"{TDOC_ID}.docx", documents["golden"]),
            ("annex/a.docx", documents["references"]),
            ("annex/b.docx", documents["branches"]),
            ("annex/c.docx", documents["golden_extra"]),
        ]
    )
    options = {
        "tdoc_id": TDOC_ID,
        "identity": IDENTITY,
        "max_documents": 100,
        "max_uncompressed_bytes": 0,
    }

    inline = convert_tdoc_zip(zip_bytes, docx_pool=None, **options)

    with DocxConvertPool(2) as pool:
        recording = _RecordingPool(pool)
        pooled = convert_tdoc_zip(zip_bytes, docx_pool=recording, **options)

    assert recording.calls == [
        (
            "markdown",
            [f"{TDOC_ID}.docx", "annex/a.docx", "annex/b.docx", "annex/c.docx"],
        )
    ]
    assert recording.worker_pids
    assert len(inline.documents) == 4
    assert inline.skipped_sources == ()
    assert pooled == inline
    assert not any(_pid_alive(pid) for pid in recording.worker_pids)

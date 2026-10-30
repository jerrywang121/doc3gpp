from __future__ import annotations

import gzip
import io
import json
import zipfile
from pathlib import Path

import numpy as np
import pytest
from sqlalchemy import text

from doc3gpp.models.search import SearchFilters
from doc3gpp.models.tdoc import TDoc
from doc3gpp.models.tdoc_content_families import store_for_type
from doc3gpp.models.tdoc_identity import document_id_for
from doc3gpp.parsers.normalizers import build_ftp_url, normalize_ftp_path
from doc3gpp.scraping.cache import TDocCache
from doc3gpp.services.semantic_search_service import SemanticSearchService
from doc3gpp.services.tdoc_family_service import PreparedFamilyExtract, TDocFamilyService
from doc3gpp.settings.loader import get_settings
from doc3gpp.storage.db.migrate import create_schema
from doc3gpp.storage.db.session import get_engine, get_session_factory
from doc3gpp.storage.repositories.search_sql import SQLAlchemySearchIndexRepository
from doc3gpp.storage.repositories.tdoc_content_sql import SQLAlchemyTDocContentRepository
from doc3gpp.storage.repositories.tdoc_sql import SQLAlchemyTDocRepository


def _make_zip(files: list[tuple[str, bytes]]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, payload in files:
            archive.writestr(name, payload)
    return output.getvalue()


def _make_docx(*paragraphs: str) -> bytes:
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _make_cr_looking_presentation() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(9), Inches(2))
    shape.text_frame.text = (
        "| CHANGE REQUEST |\n"
        "|  | 38.331 | CR | 1 | rev | 0 | Current version: | 18.5.0 |"
    )
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


class _Scraper:
    def __init__(self, payload: bytes) -> None:
        self.payload = payload
        self.calls: list[str] = []

    def get_bytes(self, url: str) -> bytes:
        self.calls.append(url)
        return self.payload


def _family_service(
    tmp_path: Path,
    payload: bytes,
    *,
    docx_pool=None,
) -> tuple[TDocFamilyService, _Scraper]:
    scraper = _Scraper(payload)
    service = TDocFamilyService(
        scraper=scraper,
        cache=TDocCache(root=tmp_path / "cache", size_limit_bytes=0),
        content_repo=SQLAlchemyTDocContentRepository(),
        max_tdoc_size_bytes=0,
        max_documents_per_zip=100,
        max_uncompressed_size_bytes=0,
        docx_pool=docx_pool,
    )
    return service, scraper


def _seed_search_context(tdoc_id: str, ftp_url: str) -> TDoc:
    create_schema()
    with get_engine().begin() as connection:
        connection.execute(
            text(
                "INSERT INTO tsgs (tsg_name, short_name, description) "
                "VALUES ('RAN WG1', 'RAN1', 'TSG RAN WG1')"
            )
        )
        connection.execute(
            text(
                """
                INSERT INTO meetings (
                    meeting_id, name, title, location, tsg, start_date,
                    end_date, ftp_url, tdoc_list_last_sync
                ) VALUES (
                    100, 'RAN1#100', 'RAN1#100', 'Online', 'RAN1',
                    '2026-01-01', '2026-01-05',
                    'https://www.3gpp.org/ftp/meetings/RAN1_100',
                    '2026-01-05T00:00:00'
                )
                """
            )
        )
    tdoc = TDoc(
        tdoc_id=tdoc_id,
        title="Presentation-bearing WID",
        meeting_id=100,
        type="WID revised",
        status="Agreed",
        ftp_url=ftp_url,
        spec="38.300",
        release="Rel-18",
    )
    SQLAlchemyTDocRepository().upsert(tdoc)
    return tdoc


def _prepare_family(
    tmp_path: Path,
    tdoc: TDoc,
    payload: bytes,
    *,
    docx_pool=None,
) -> tuple[TDocFamilyService, object, str]:
    create_schema()
    store = store_for_type(tdoc.type)
    assert store is not None
    SQLAlchemyTDocRepository().upsert(tdoc)
    service, _scraper = _family_service(tmp_path, payload, docx_pool=docx_pool)
    prepared = service.prepare_extract(tdoc, store)
    assert isinstance(prepared, PreparedFamilyExtract)
    canonical = normalize_ftp_path(build_ftp_url(tdoc.ftp_url or ""))
    return service, prepared, canonical


def _assert_gzip_json_content(table: str, document_id: str, markdown: str) -> None:
    with get_session_factory()() as session:
        compressed = session.execute(
            text(f"SELECT content FROM {table} WHERE document_id = :document_id"),
            {"document_id": document_id},
        ).scalar_one()
    assert bytes(compressed).startswith(b"\x1f\x8b")
    assert json.loads(gzip.decompress(bytes(compressed))) == markdown


def test_mixed_family_presentation_is_compressed_and_fts_searchable(
    sqlite_env, tmp_path: Path,
) -> None:
    from tests.fixtures.presentation_builders import build_presentation_with_content

    tdoc_id = "R5-260301"
    ftp_url = f"r5/presentation/{tdoc_id}.zip"
    tdoc = _seed_search_context(tdoc_id, ftp_url)
    presentation_source = "slides/presentation.pptx"
    payload = _make_zip(
        [
            (presentation_source, build_presentation_with_content()),
            (f"{tdoc_id}.docx", _make_docx("family main body")),
        ]
    )

    service, prepared, canonical = _prepare_family(tmp_path, tdoc, payload)
    assert [document.source_file for document in prepared.documents] == [
        f"{tdoc_id}.docx",
        presentation_source,
    ]
    service.persist_prepared_extract(prepared)

    store = store_for_type(tdoc.type)
    assert store is not None
    stored = SQLAlchemyTDocContentRepository().list_documents(store, canonical)
    assert [document.source_file for document in stored] == [
        f"{tdoc_id}.docx",
        presentation_source,
    ]
    assert [document.document_id for document in stored] == [
        document_id_for(canonical, f"{tdoc_id}.docx"),
        document_id_for(canonical, presentation_source),
    ]
    presentation = stored[1]
    assert "second slide text" in presentation.markdown_text
    assert "speaker secret" in presentation.markdown_text
    _assert_gzip_json_content(
        "tdoc_wid_contents",
        presentation.document_id,
        presentation.markdown_text,
    )

    fts = SQLAlchemySearchIndexRepository()
    fts.upsert(tdoc_id)
    for query in ("second slide text", "speaker secret"):
        hits = fts.search(query, SearchFilters(limit=10))
        assert [hit.tdoc_id for hit in hits] == [tdoc_id]
        assert any(
            match.source_file == presentation_source
            for match in hits[0].matched_documents
        )


def test_empty_presentation_only_family_zip_is_stored_as_a_known_document(
    sqlite_env, tmp_path: Path,
) -> None:
    from tests.fixtures.presentation_builders import build_empty_presentation

    tdoc_id = "R5-260302"
    ftp_url = f"r5/presentation/{tdoc_id}.zip"
    tdoc = TDoc(tdoc_id=tdoc_id, type="WID new", ftp_url=ftp_url)
    payload = _make_zip([(f"{tdoc_id}.pptx", build_empty_presentation())])

    service, prepared, canonical = _prepare_family(tmp_path, tdoc, payload)
    assert [(document.source_file, document.markdown_text) for document in prepared.documents] == [
        (f"{tdoc_id}.pptx", ""),
    ]
    service.persist_prepared_extract(prepared)

    store = store_for_type(tdoc.type)
    assert store is not None
    stored = SQLAlchemyTDocContentRepository().list_documents(store, canonical)
    assert [(document.source_file, document.markdown_text) for document in stored] == [
        (f"{tdoc_id}.pptx", ""),
    ]
    _assert_gzip_json_content("tdoc_wid_contents", stored[0].document_id, "")


def test_incomplete_family_presentation_annex_is_retried_on_normal_parse(
    sqlite_env, tmp_path: Path,
) -> None:
    tdoc_id = "R5-260303"
    ftp_url = f"r5/presentation/{tdoc_id}.zip"
    tdoc = TDoc(tdoc_id=tdoc_id, type="WID revised", ftp_url=ftp_url)
    payload = _make_zip(
        [
            (f"{tdoc_id}.docx", b"word"),
            ("slides/retried.pptx", b"presentation"),
        ]
    )

    class _RetryPool:
        def __init__(self) -> None:
            self.calls = 0

        def convert_members(self, kind, members):
            assert kind == "markdown"
            self.calls += 1
            outcomes = []
            for source_file, _payload in members:
                if source_file.endswith(".pptx") and self.calls == 1:
                    outcomes.append(RuntimeError("presentation annex failed"))
                elif source_file.endswith(".pptx"):
                    outcomes.append("retried slide text")
                else:
                    outcomes.append("family main text")
            return outcomes

    create_schema()
    SQLAlchemyTDocRepository().upsert(tdoc)
    store = store_for_type(tdoc.type)
    assert store is not None
    pool = _RetryPool()
    service, _scraper = _family_service(tmp_path, payload, docx_pool=pool)

    first = service.prepare_extract(tdoc, store)
    assert isinstance(first, PreparedFamilyExtract)
    assert first.skipped_sources == ("slides/retried.pptx",)
    service.persist_prepared_extract(first)

    second = service.prepare_extract(tdoc, store)
    assert isinstance(second, PreparedFamilyExtract)
    assert pool.calls == 2
    assert second.skipped_sources == ()
    assert [document.source_file for document in second.documents] == [
        f"{tdoc_id}.docx",
        "slides/retried.pptx",
    ]
    service.persist_prepared_extract(second)

    canonical = normalize_ftp_path(build_ftp_url(ftp_url))
    meta = SQLAlchemyTDocContentRepository().get_extract_meta(canonical)
    assert meta is not None
    assert meta.skipped_sources == ()


def test_cr_presentation_is_a_compressed_sql_annex_not_the_cover(tmp_path: Path, sqlite_env) -> None:
    from doc3gpp.parsers.cr.header import is_cr_header_present
    from doc3gpp.parsers.docx_converter import convert_document_to_markdown
    from doc3gpp.scraping.cache import TDocCache
    from doc3gpp.services.tdoc_cr_service import TDocCrService
    from doc3gpp.storage.repositories.tdoc_cr_change_details_sql import (
        SQLAlchemyTDocCrChangeDetailsRepository,
    )
    from doc3gpp.storage.repositories.tdoc_cr_sql import SQLAlchemyTDocCrRepository
    from doc3gpp.storage.repositories.tdoc_cr_ttcn_sql import (
        SQLAlchemyTDocCrTtcnRepository,
    )
    from doc3gpp.storage.repositories.tdoc_zip_inventory_sql import (
        SQLAlchemyTDocZipInventoryRepository,
    )
    tdoc_id = "R5s260009"
    ftp_url = f"tsg_ran/wg5_test_ex-t1/ttcn/ttcn_crs/2026/docs/{tdoc_id}.zip"
    source_file = "CRs/presentation.pptx"
    cr_presentation = _make_cr_looking_presentation()
    cr_markdown = convert_document_to_markdown(cr_presentation, source_file)
    assert is_cr_header_present(cr_markdown)[0]

    fixture_path = Path(__file__).parent.parent / "fixtures" / "tdoc_cr_doc" / f"{tdoc_id}.zip"
    with zipfile.ZipFile(fixture_path) as source_archive:
        members = [
            (info.filename, source_archive.read(info))
            for info in source_archive.infolist()
            if not info.is_dir()
        ]
    payload = _make_zip([*members, (source_file, cr_presentation)])
    create_schema()
    tdoc = TDoc(tdoc_id=tdoc_id, type="CR", ftp_url=ftp_url)
    SQLAlchemyTDocRepository().upsert(tdoc)
    service = TDocCrService(
        cache=TDocCache(root=tmp_path / "cr-cache", size_limit_bytes=0),
        scraper_client=_Scraper(payload),
        cr_repository=SQLAlchemyTDocCrRepository(),
        cr_ttcn_repository=SQLAlchemyTDocCrTtcnRepository(),
        cr_change_details_repository=SQLAlchemyTDocCrChangeDetailsRepository(),
        tdoc_repository=SQLAlchemyTDocRepository(),
        content_repo=SQLAlchemyTDocContentRepository(),
        inventory_repository=SQLAlchemyTDocZipInventoryRepository(),
    )

    result = service.extract(tdoc_id)

    assert result.extract_meta.doc_filename.lower().endswith((".docx", ".doc"))
    assert result.extract_meta.doc_filename != source_file
    store = store_for_type("CR")
    assert store is not None
    canonical = normalize_ftp_path(ftp_url)
    annexes = SQLAlchemyTDocContentRepository().list_documents(store, canonical)
    assert source_file in [document.source_file for document in annexes]
    presentation = next(
        document for document in annexes if document.source_file == source_file
    )
    assert "CHANGE REQUEST" in presentation.markdown_text
    _assert_gzip_json_content(
        "tdoc_cr_annexes",
        presentation.document_id,
        presentation.markdown_text,
    )


@pytest.mark.semantic
def test_presentation_slide_and_notes_text_are_found_by_semantic_search(
    sqlite_env, tmp_path: Path,
) -> None:
    from tests.fixtures.presentation_builders import build_presentation_with_content

    tdoc_id = "R5-260304"
    ftp_url = f"r5/presentation/{tdoc_id}.zip"
    tdoc = _seed_search_context(tdoc_id, ftp_url)
    presentation_source = "slides/presentation.pptx"
    payload = _make_zip(
        [
            (presentation_source, build_presentation_with_content()),
            (f"{tdoc_id}.docx", _make_docx("family main body")),
        ]
    )
    service, prepared, canonical = _prepare_family(tmp_path, tdoc, payload)
    service.persist_prepared_extract(prepared)

    fts = SQLAlchemySearchIndexRepository()
    fts.upsert(tdoc_id)
    queries = ("second slide text", "speaker secret")
    presentation_id = document_id_for(canonical, presentation_source)
    for query in queries:
        fts_hits = fts.search(query, SearchFilters(limit=5))
        assert [hit.tdoc_id for hit in fts_hits] == [tdoc_id]
        assert any(
            match.source_file == presentation_source
            for match in fts_hits[0].matched_documents
        )

    from doc3gpp.storage.repositories.vector_sql import SQLAlchemyVectorIndexRepository

    vector_repo = SQLAlchemyVectorIndexRepository()

    class _Embedder:
        def __init__(self, dimension: int) -> None:
            self.dim = dimension

        def encode(self, texts: list[str]) -> np.ndarray:
            vectors = np.zeros((len(texts), self.dim), dtype=np.float32)
            for row, value in enumerate(texts):
                lowered = value.casefold()
                slide_text = "second slide text" in lowered
                notes_text = "speaker secret" in lowered
                if slide_text:
                    vectors[row, 0] = 1.0
                if notes_text:
                    vectors[row, 1] = 1.0
                if not slide_text and not notes_text:
                    vectors[row, 2] = 1.0
            return vectors

    assert vector_repo._dim > 2
    semantic = SemanticSearchService(
        fts,
        _Embedder(vector_repo._dim),
        vector_repo,
        get_settings(),
    )
    semantic.index_for_tdoc(tdoc_id)

    for query in queries:
        hits = semantic.search(
            query,
            fts5_query=None,
            filters=SearchFilters(limit=5),
            limit=5,
            fts5_weight=0.5,
        )
        assert [hit.tdoc_id for hit in hits] == [tdoc_id]
        assert any(
            document.document_id == presentation_id
            for hit in hits
            for document in hit.document_hits
        )


def test_mixed_family_presentation_details_survive_cache_eviction(
    sqlite_env, tmp_path: Path, monkeypatch,
) -> None:
    import asyncio
    import json
    import shutil

    from fastapi.testclient import TestClient
    from typer.testing import CliRunner

    from doc3gpp.cli import app as cli_app
    from doc3gpp.models.tdoc_content_families import store_for_type
    from doc3gpp.services.tdoc_content_service import TDocContentService
    from doc3gpp.settings.schema import (
        CacheSettings,
        MCPSettings,
        ServerSettings,
        Settings,
    )
    from doc3gpp.web.app import build_app, build_state
    from doc3gpp.web.mcp_server import build_mcp_server
    from tests.fixtures.presentation_builders import build_presentation_with_content

    tdoc_id = "R5-260305"
    ftp_url = f"r5/presentation/{tdoc_id}.zip"
    source_file = "slides/presentation.pptx"
    tdoc = _seed_search_context(tdoc_id, ftp_url)
    payload = _make_zip(
        [
            (f"{tdoc_id}.docx", _make_docx("family main body")),
            (source_file, build_presentation_with_content()),
        ]
    )
    service, _scraper = _family_service(tmp_path, payload)
    store = store_for_type(tdoc.type)
    assert store is not None
    prepared = service.prepare_extract(tdoc, store)
    assert isinstance(prepared, PreparedFamilyExtract)
    service.persist_prepared_extract(prepared)
    expected_markdown = next(
        document.markdown_text
        for document in prepared.documents
        if document.source_file == source_file
    )

    cache_root = tmp_path / "cache"
    shutil.rmtree(cache_root)
    assert not cache_root.exists()

    def no_cache_read(*_args, **_kwargs):
        pytest.fail("stored family presentation reads must not use the Markdown cache")

    monkeypatch.setattr(TDocContentService, "_read_cache_archive", no_cache_read)
    monkeypatch.setattr(
        "doc3gpp.cli.trigger_auto_sync", lambda **_kwargs: None
    )
    monkeypatch.setattr(
        "doc3gpp.cli.build_tdoc_cr_service",
        lambda *_args, **_kwargs: pytest.fail(
            "a complete selected presentation must not trigger extraction"
        ),
    )

    cli = CliRunner().invoke(
        cli_app,
        [
            "tdoc",
            "show",
            "--tdoc",
            tdoc_id,
            "--format",
            "raw",
            "--source-file",
            source_file,
        ],
    )
    assert cli.exit_code == 0, cli.output
    assert cli.stdout == expected_markdown + ("" if expected_markdown.endswith("\n") else "\n")

    settings = Settings(
        server=ServerSettings(enabled=True, port=8765),
        mcp=MCPSettings(enabled=True),
        cache=CacheSettings(dir=cache_root),
    )
    with TestClient(build_app(settings)) as client:
        http = client.get(
            f"/tdocs/{tdoc_id}/content",
            params={"source_file": source_file, "format": "json"},
        )
    assert http.status_code == 200, http.text
    http_json = http.json()
    assert http_json["complete"] is True
    assert http_json["documents"][0]["source_file"] == source_file
    assert http_json["documents"][0]["markdown"] == expected_markdown

    state = build_state(settings)
    try:
        server = build_mcp_server(state)
        result = asyncio.run(
            server.call_tool(
                "get_tdoc_content",
                {
                    "tdoc_id": tdoc_id,
                    "source_file": source_file,
                    "format": "json",
                },
            )
        )
        assert result.is_error is False
        mcp_json = result.content[0].text
        assert json.loads(mcp_json)["documents"][0]["markdown"] == expected_markdown
    finally:
        if state.close_docx_pool is not None:
            state.close_docx_pool()
        if state.close_embedder is not None:
            state.close_embedder()
        state.engine.dispose()
        state.testcase_engine.dispose()
        if state.specdata_engine is not None:
            state.specdata_engine.dispose()

    assert http.content.decode("utf-8") == mcp_json
    assert not cache_root.exists()


def test_force_reparse_refreshes_presentation_without_changing_word_main(
    sqlite_env, tmp_path: Path,
) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    tdoc_id = "R5-260306"
    ftp_url = f"r5/presentation/{tdoc_id}.zip"
    tdoc = _seed_search_context(tdoc_id, ftp_url)
    store = store_for_type(tdoc.type)
    assert store is not None
    source_file = "slides/presentation.pptx"

    def mixed_zip(slide_text: str, notes_text: str) -> bytes:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(1)
        )
        shape.text_frame.text = slide_text
        slide.notes_slide.notes_text_frame.text = notes_text
        output = io.BytesIO()
        presentation.save(output)
        return _make_zip(
            [
                (f"{tdoc_id}.docx", _make_docx("unchanged Word main")),
                (source_file, output.getvalue()),
            ]
        )

    initial_service, _initial_scraper = _family_service(
        tmp_path / "initial", mixed_zip("initial slide signal", "initial notes")
    )
    initial = initial_service.prepare_extract(tdoc, store)
    assert isinstance(initial, PreparedFamilyExtract)
    initial_service.persist_prepared_extract(initial)
    canonical = normalize_ftp_path(build_ftp_url(ftp_url))
    before = SQLAlchemyTDocContentRepository().list_documents(store, canonical)
    old_main = before[0].markdown_text
    old_presentation = before[1].markdown_text

    forced_service, forced_scraper = _family_service(
        tmp_path / "forced", mixed_zip("refreshed slide signal", "refreshed notes")
    )
    refreshed = forced_service.prepare_extract(tdoc, store, force=True)
    assert isinstance(refreshed, PreparedFamilyExtract)
    forced_service.persist_prepared_extract(refreshed)
    after = SQLAlchemyTDocContentRepository().list_documents(store, canonical)

    assert forced_scraper.calls
    assert after[0].source_file == f"{tdoc_id}.docx"
    assert after[0].markdown_text == old_main
    assert after[1].source_file == source_file
    assert after[1].markdown_text != old_presentation
    assert "refreshed slide signal" in after[1].markdown_text
    assert "refreshed notes" in after[1].markdown_text


def test_legacy_ppt_family_member_uses_impress_or_deterministic_converter(
    sqlite_env, tmp_path: Path, monkeypatch,
) -> None:
    import shutil
    import subprocess

    from doc3gpp.parsers import presentation_converter
    from tests.fixtures.presentation_builders import build_presentation_with_content

    tdoc_id = "R5-260307"
    ftp_url = f"r5/presentation/{tdoc_id}.zip"
    tdoc = _seed_search_context(tdoc_id, ftp_url)
    source_file = "slides/legacy.ppt"
    pptx_bytes = build_presentation_with_content()
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        monkeypatch.setattr(
            presentation_converter,
            "_as_pptx_bytes",
            lambda _payload, _filename: pptx_bytes,
        )
        legacy_ppt = b"deterministic legacy PPT fixture"
    else:
        source_dir = tmp_path / "pptx-source"
        output_dir = tmp_path / "legacy-ppt"
        source_dir.mkdir()
        output_dir.mkdir()
        source = source_dir / "presentation.pptx"
        source.write_bytes(pptx_bytes)
        result = subprocess.run(
            [
                executable,
                f"-env:UserInstallation={(tmp_path / 'legacy-export-profile').as_uri()}",
                "--headless",
                "--nologo",
                "--nodefault",
                "--norestore",
                "--convert-to",
                "ppt",
                "--outdir",
                str(output_dir),
                str(source),
            ],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
        assert result.returncode == 0, result.stderr or result.stdout
        legacy_path = output_dir / "presentation.ppt"
        assert legacy_path.is_file(), result.stderr or result.stdout
        legacy_ppt = legacy_path.read_bytes()

    payload = _make_zip(
        [
            (f"{tdoc_id}.docx", _make_docx("family main body")),
            (source_file, legacy_ppt),
        ]
    )
    tdoc_service, prepared, canonical = _prepare_family(tmp_path, tdoc, payload)
    tdoc_service.persist_prepared_extract(prepared)
    store = store_for_type(tdoc.type)
    assert store is not None
    stored = SQLAlchemyTDocContentRepository().list_documents(store, canonical)
    presentation = next(
        document for document in stored if document.source_file == source_file
    )

    assert presentation.source_file == source_file
    assert "second slide text" in presentation.markdown_text
    assert "speaker secret" in presentation.markdown_text

"""Synthetic ``.docx`` builders for the DOCX converter golden tests.

:func:`build_golden_docx` returns one document that exercises the shapes
the converter renders: heading levels, title and list styles, run
formatting, a captioned table with horizontal and vertical merges and a
``|`` in a cell, tracked insertions and deletions, and an unstyled
paragraph.
"""

from __future__ import annotations

import io

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _text_run(text: str, text_tag: str = "w:t"):
    run = OxmlElement("w:r")
    node = OxmlElement(text_tag)
    node.text = text
    run.append(node)
    return run


def _revision(tag: str, revision_id: int, run):
    element = OxmlElement(tag)
    element.set(qn("w:id"), str(revision_id))
    element.set(qn("w:author"), "Golden Fixture")
    element.set(qn("w:date"), "2026-01-01T00:00:00Z")
    element.append(run)
    return element


def _add_tracked_paragraph(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph._p.append(_text_run("Kept text, "))
    paragraph._p.append(_revision("w:ins", 1, _text_run("inserted text")))
    paragraph._p.append(_text_run(" and "))
    paragraph._p.append(_revision("w:del", 2, _text_run("deleted text", "w:delText")))
    paragraph._p.append(_text_run(" end."))


def _add_captioned_table(document: Document) -> None:
    document.add_paragraph("Table 5.1-1: Example caption")
    table = document.add_table(rows=3, cols=3)
    for col, text in enumerate(["Header A", "Header B", "Header C"]):
        table.cell(0, col).text = text
    table.cell(1, 0).text = "Spanned body"
    table.cell(1, 2).text = "Vertical top"
    table.cell(2, 0).text = "R2C0"
    table.cell(2, 1).text = "a | b"
    table.cell(1, 0).merge(table.cell(1, 1))
    table.cell(1, 2).merge(table.cell(2, 2))


def build_golden_docx() -> bytes:
    document = Document()
    document.add_paragraph("Golden Characterisation", style="Title")
    document.add_paragraph("1 Scope", style="Heading 1")

    formatted = document.add_paragraph("Plain ")
    formatted.add_run("bold").bold = True
    formatted.add_run(" ")
    formatted.add_run("italic").italic = True
    formatted.add_run(" ")
    bold_italic = formatted.add_run("both")
    bold_italic.bold = True
    bold_italic.italic = True

    document.add_paragraph("5.1 General", style="Heading 2")
    document.add_paragraph("First bullet item", style="List Bullet")
    _add_captioned_table(document)
    _add_tracked_paragraph(document)
    document.add_paragraph("Closing paragraph without a style.")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()

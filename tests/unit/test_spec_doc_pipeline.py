import asyncio
import io
import threading
import zipfile
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import MagicMock

from docx import Document

from doc3gpp.models.jobs import Job, JobKind, JobStatus
from doc3gpp.models.spec import SpecVersion
from doc3gpp.models.spec_doc import SpecDocSource
from doc3gpp.models.spec_doc_zip_inventory import (
    SpecDocInventorySchemaError,
    SpecDocZipInventory,
)
from doc3gpp.services.pipeline_runtime import PipelineRuntime
from doc3gpp.services.spec_doc_service import SpecDocService
from doc3gpp.settings.schema import Settings
from doc3gpp.web.workers.handlers import _parse_spec_docs


def _zip_bytes() -> bytes:
    doc = Document()
    doc.add_heading("1 Scope", level=1)
    doc.add_paragraph("prepared body")
    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as archive:
        archive.writestr("38331-j30.docx", doc_buffer.getvalue())
    return zip_buffer.getvalue()


class _MemoryInventoryRepo:
    def __init__(self):
        self.items = {}
        self.upserts = []
        self.write_threads = []
        self.error = None

    def upsert(self, inventory):
        self.upserts.append(inventory)
        self.write_threads.append(threading.get_ident())
        if self.error is not None:
            raise self.error
        self.items[(inventory.spec_id, inventory.version)] = inventory

    def get(self, spec_id, version):
        return self.items.get((spec_id, version))


class _FailingDocxPool:
    def __init__(self, error):
        self.error = error

    def convert_members(self, _kind, members):
        return [self.error for _ in members]


def _service(repo, tmp_path, inventory_repo=None, **kwargs):
    spec_repo = MagicMock()
    spec_repo.list_versions.return_value = [
        SpecVersion("38.331", "18.5.0", "https://example.test/doc.zip", "Rel-18")
    ]
    return SpecDocService(
        spec_repo=spec_repo,
        doc_repo=repo,
        inventory_repo=(
            inventory_repo if inventory_repo is not None else MagicMock()
        ),
        fetcher=lambda _: _zip_bytes(),
        cache_dir=tmp_path / "spec-cache",
        **kwargs,
    )


def test_prepare_parse_does_not_mutate_repository(tmp_path):
    repo = MagicMock()
    inventory_repo = MagicMock()
    repo.get_source.return_value = None
    service = _service(repo, tmp_path, inventory_repo=inventory_repo)

    prepared = service.prepare_parse("38.331")

    assert prepared.source.version == "18.5.0"
    assert prepared.toc.entries[0].section_no == "1"
    assert prepared.drafts
    assert prepared.inventory == SpecDocZipInventory(
        "38.331", "18.5.0", "Rel-18", ("38331-j30.docx",)
    )
    inventory_repo.upsert.assert_not_called()
    repo.record_download.assert_not_called()
    repo.upsert_toc.assert_not_called()
    repo.replace_chunks.assert_not_called()
    repo.record_parsed.assert_not_called()


def test_persist_prepared_parse_writes_and_returns_source(tmp_path):
    repo = MagicMock()
    inventory_repo = MagicMock()
    source = SpecDocSource("38.331", "18.5.0", parsed_at=object())
    repo.get_source.side_effect = [None, source]
    events = []
    inventory_repo.upsert.side_effect = lambda _inventory: events.append("inventory")
    repo.record_download.side_effect = lambda *_args, **_kwargs: events.append("download")
    service = _service(repo, tmp_path, inventory_repo=inventory_repo)

    prepared = service.prepare_parse("38.331")
    result = service.persist_prepared_parse(prepared)

    assert result is source
    inventory_repo.upsert.assert_called_once_with(prepared.inventory)
    assert events[0] == "inventory"
    repo.record_download.assert_called_once()
    repo.upsert_toc.assert_called_once()
    repo.replace_chunks.assert_called_once()
    repo.record_parsed.assert_called_once_with("38.331", "18.5.0", chunk_count=1)


def test_queued_parse_uses_and_releases_one_shared_permit():
    pipeline = PipelineRuntime(prepare_workers=1, max_pending=1)
    prepared = object()
    source = SpecDocSource("38.331", "18.5.0", parsed_at=object())
    calls: list[str] = []

    class FakeSpecDocService:
        def prepare_parse(self, spec_id, **kwargs):
            assert spec_id == "38.331"
            calls.append("prepare")
            return prepared

        def persist_prepared_parse(self, value):
            assert value is prepared
            calls.append("persist")
            return source

    job = Job(
        id="spec-doc-job",
        kind=JobKind.PARSE_SPEC_DOCS,
        status=JobStatus.RUNNING,
        params={"spec_id": "38.331"},
        log_lines=(),
        result_summary=None,
        error=None,
        created_at=datetime.now(timezone.utc),
        started_at=None,
        finished_at=None,
    )
    services = SimpleNamespace(pipeline=pipeline, spec_doc=FakeSpecDocService())

    async def run():
        return await _parse_spec_docs(
            job,
            services,
            Settings(),
            progress=lambda *_args, **_kwargs: None,
            cancel_event=asyncio.Event(),
        )

    try:
        result = asyncio.run(run())

        async def reserve_after_handler():
            permit = await asyncio.wait_for(pipeline.reserve_async(), timeout=1)
            permit.release()

        asyncio.run(reserve_after_handler())
    finally:
        pipeline.shutdown()

    assert calls == ["prepare", "persist"]
    assert result == {
        "requested": 1,
        "successes": 1,
        "skipped": 0,
        "failures": 0,
        "failure_details": {},
    }


def test_persist_records_source_when_zip_cached_but_source_row_missing(tmp_path):
    repo = MagicMock()
    source = SpecDocSource("38.331", "18.5.0", parsed_at=object())
    repo.get_source.side_effect = [None, source]
    service = _service(repo, tmp_path)
    service._write_zip_cache("38.331", "18.5.0", _zip_bytes())
    service._fetcher = MagicMock(side_effect=AssertionError("must use cache"))

    prepared = service.prepare_parse("38.331")
    assert prepared.download_required is False
    service.persist_prepared_parse(prepared)

    repo.record_download.assert_called_once_with(
        "38.331",
        "18.5.0",
        release="Rel-18",
        ftp_url="https://example.test/doc.zip",
        docx_count=1,
    )
    names = [c[0] for c in repo.method_calls]
    assert names.index("record_download") < names.index("record_parsed")


def test_persist_keeps_existing_source_metadata_when_cached(tmp_path):
    repo = MagicMock()
    existing = SpecDocSource("38.331", "18.5.0", parsed_at=None)
    repo.get_source.side_effect = [existing, existing]
    service = _service(repo, tmp_path)
    service._write_zip_cache("38.331", "18.5.0", _zip_bytes())

    service.persist_prepared_parse(service.prepare_parse("38.331"))

    repo.record_download.assert_not_called()
    repo.record_parsed.assert_called_once()


def _job(params):
    return Job(
        id="spec-doc-job",
        kind=JobKind.PARSE_SPEC_DOCS,
        status=JobStatus.RUNNING,
        params=params,
        log_lines=(),
        result_summary=None,
        error=None,
        created_at=datetime.now(timezone.utc),
        started_at=None,
        finished_at=None,
    )


def _run_handler(
    spec_service, params=None, *, pipeline=None, cancel_event=None
):
    owns_pipeline = pipeline is None
    if pipeline is None:
        pipeline = PipelineRuntime(prepare_workers=1, max_pending=1)
    messages: list[str] = []
    services = SimpleNamespace(pipeline=pipeline, spec_doc=spec_service)
    if cancel_event is None:
        cancel_event = asyncio.Event()
    try:
        result = asyncio.run(
            _parse_spec_docs(
                _job(params or {"spec_id": "38.331"}),
                services,
                Settings(),
                progress=lambda msg, **_kw: messages.append(msg),
                cancel_event=cancel_event,
            )
        )
    finally:
        if owns_pipeline:
            pipeline.shutdown()
    return result, messages


def test_queued_parse_already_parsed_is_skipped():
    class Svc:
        def prepare_parse(self, spec_id, **kwargs):
            return SpecDocSource("38.331", "18.5.0", parsed_at=object())

        def persist_prepared_parse(self, value):
            raise AssertionError("must not persist")

    result, messages = _run_handler(Svc())
    assert result["skipped"] == 1 and result["successes"] == 0
    assert result["failures"] == 0
    assert not any("done" in m and "@" in m for m in messages)


def test_queued_parse_oversized_is_skipped():
    from doc3gpp.models.spec_doc import SpecDocTooLargeError

    class Svc:
        def prepare_parse(self, spec_id, **kwargs):
            raise SpecDocTooLargeError("too big")

    result, _ = _run_handler(Svc())
    assert result["skipped"] == 1
    assert result["failures"] == 0
    assert result["failure_details"] == {}


def test_queued_parse_failure_reason_is_exact():
    class Svc:
        def prepare_parse(self, spec_id, **kwargs):
            raise RuntimeError("boom happened")

    result, messages = _run_handler(Svc())
    assert result["failure_details"] == {"38.331": "boom happened"}
    assert "spec 38.331 failed: boom happened" in messages


def test_queued_parse_success_emits_spec_done_progress():
    source = SpecDocSource("38.331", "18.5.0", parsed_at=object())

    class Svc:
        def prepare_parse(self, spec_id, **kwargs):
            return object()

        def persist_prepared_parse(self, value):
            return source

    result, messages = _run_handler(Svc())
    assert result["successes"] == 1
    assert "spec 38.331@18.5.0 done" in messages


def test_failed_preparation_and_failed_inventory_write_preserves_original(
    tmp_path, caplog
):
    repo = MagicMock()
    repo.get_source.return_value = None
    inventory_repo = _MemoryInventoryRepo()
    inventory_repo.error = OSError("inventory write failed")
    original = RuntimeError("conversion exploded")
    service = _service(
        repo,
        tmp_path,
        inventory_repo=inventory_repo,
        docx_pool=_FailingDocxPool(original),
    )

    result = service.parse_many(["38.331"])

    assert result.failures["38.331"] == "conversion exploded"
    assert "inventory" in caplog.text.lower()
    assert len(inventory_repo.upserts) == 1
    repo.record_download.assert_not_called()
    repo.upsert_toc.assert_not_called()
    repo.replace_chunks.assert_not_called()
    repo.record_parsed.assert_not_called()


def test_successful_parse_inventory_schema_failure_does_not_create_parsed_content(
    tmp_path
):
    repo = MagicMock()
    repo.get_source.return_value = None
    inventory_repo = _MemoryInventoryRepo()
    inventory_repo.error = SpecDocInventorySchemaError(
        "run 'doc3gpp db init --scope specdata' to upgrade"
    )
    service = _service(repo, tmp_path, inventory_repo=inventory_repo)

    result = service.parse_many(["38.331"])

    assert result.failures["38.331"] == (
        "run 'doc3gpp db init --scope specdata' to upgrade"
    )
    assert len(inventory_repo.upserts) == 1
    repo.record_download.assert_not_called()
    repo.upsert_toc.assert_not_called()
    repo.replace_chunks.assert_not_called()
    repo.record_parsed.assert_not_called()


def test_queued_failure_persists_on_writer_before_reporting_original(tmp_path):
    repo = MagicMock()
    existing = SpecDocSource("38.331", "18.5.0", parsed_at=object())
    repo.get_source.return_value = existing
    inventory_repo = _MemoryInventoryRepo()
    original = RuntimeError("conversion exploded")
    service = _service(
        repo,
        tmp_path,
        inventory_repo=inventory_repo,
        docx_pool=_FailingDocxPool(original),
    )
    prepare_threads = []
    writer_threads = []
    original_prepare = service.prepare_parse
    original_persist = service.persist_prepared_parse

    def record_prepare_thread(*args, **kwargs):
        prepare_threads.append(threading.get_ident())
        return original_prepare(*args, **kwargs)

    def record_writer_thread(prepared):
        writer_threads.append(threading.get_ident())
        return original_persist(prepared)

    service.prepare_parse = record_prepare_thread
    service.persist_prepared_parse = record_writer_thread
    pipeline = PipelineRuntime(prepare_workers=1, max_pending=1)
    try:
        result, messages = _run_handler(
            service,
            {"spec_id": "38.331", "force": True},
            pipeline=pipeline,
        )

        async def reserve_after_handler():
            permit = await asyncio.wait_for(pipeline.reserve_async(), timeout=1)
            permit.release()

        asyncio.run(reserve_after_handler())
    finally:
        pipeline.shutdown()

    assert result["failure_details"] == {"38.331": "conversion exploded"}
    assert "spec 38.331 failed: conversion exploded" in messages
    assert prepare_threads and writer_threads and inventory_repo.write_threads
    assert inventory_repo.write_threads[0] == writer_threads[0]
    assert prepare_threads[0] != writer_threads[0]
    assert inventory_repo.get("38.331", "18.5.0").zip_files == (
        "38331-j30.docx",
    )
    repo.invalidate_parse_state.assert_not_called()
    repo.record_parsed.assert_not_called()
    assert existing.parsed_at is not None


def test_queued_cancel_after_preparation_prevents_inventory_write(tmp_path):
    repo = MagicMock()
    repo.get_source.return_value = None
    inventory_repo = _MemoryInventoryRepo()
    service = _service(repo, tmp_path, inventory_repo=inventory_repo)
    cancel_event = threading.Event()
    original_prepare = service.prepare_parse
    original_persist = service.persist_prepared_parse
    persist_calls = []

    def prepare_then_cancel(*args, **kwargs):
        prepared = original_prepare(*args, **kwargs)
        cancel_event.set()
        return prepared

    def record_persist(prepared):
        persist_calls.append(prepared)
        return original_persist(prepared)

    service.prepare_parse = prepare_then_cancel
    service.persist_prepared_parse = record_persist
    pipeline = PipelineRuntime(prepare_workers=1, max_pending=1)
    try:
        try:
            _run_handler(
                service,
                pipeline=pipeline,
                cancel_event=cancel_event,
            )
        except asyncio.CancelledError:
            pass
        else:
            raise AssertionError("queued parse should observe cancellation before persist")

        async def reserve_after_cancellation():
            permit = await asyncio.wait_for(pipeline.reserve_async(), timeout=1)
            permit.release()

        asyncio.run(reserve_after_cancellation())
    finally:
        pipeline.shutdown()

    assert persist_calls == []
    assert inventory_repo.upserts == []
    repo.record_parsed.assert_not_called()

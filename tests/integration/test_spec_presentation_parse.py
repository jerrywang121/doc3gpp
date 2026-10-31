"""Integration tests for searchable spec-document presentations."""

from __future__ import annotations

import io
import zipfile
from itertools import pairwise

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from doc3gpp.models.spec import SpecVersion
from doc3gpp.models.spec_doc import SpecDocError, SpecDocSearchFilters
from doc3gpp.services.spec_doc_search_service import SpecDocSearchService
from doc3gpp.services.spec_doc_semantic_service import SpecDocSemanticService
from doc3gpp.services.spec_doc_service import SpecDocService
from doc3gpp.settings.loader import get_settings
from doc3gpp.storage.db.migrate import create_schema
from doc3gpp.storage.repositories.spec_doc_sql import SQLAlchemySpecDocRepository
from doc3gpp.storage.repositories.spec_doc_zip_inventory_sql import (
    SQLAlchemySpecDocZipInventoryRepository,
)
from tests.fixtures.presentation_builders import build_empty_presentation

_SPEC_ID = "38.331"
_VERSION = "18.5.0"
_FTP_URL = "https://example.test/38.331.zip"
_LONG_SLIDE_TEXT = (
    "Presentation data is intentionally long enough to split into bounded search "
    "chunks without losing the full source Markdown."
)
_EXPECTED_MARKDOWN = (
    "## Slide 1\n\n"
    "Contents\n\n"
    "1 Scope\n\n"
    "Table 1: Equipment\n\n"
    "slideonlysignal\n\n"
    f"{_LONG_SLIDE_TEXT}\n\n"
    "| Equipment | Status |\n"
    "| --- | --- |\n"
    "| NR-ARFCN | supported |\n\n"
    "**Speaker notes:**\n\n"
    "speaker secret notesonlyprobe"
)


class _SpecRepo:
    def list_versions(self, spec_id: str, **_filters) -> list[SpecVersion]:
        if spec_id != _SPEC_ID:
            return []
        return [SpecVersion(_SPEC_ID, _VERSION, _FTP_URL, "Rel-18")]


class _RecordingEmbedder:
    def __init__(self) -> None:
        self.texts: list[str] = []

    def encode(self, texts: list[str]) -> list[list[float]]:
        self.texts.extend(texts)
        return [[float(index)] for index, _text in enumerate(texts)]


class _RecordingVectorRepository:
    def __init__(self) -> None:
        self.upserts: list[tuple[str, str, list[list[float]]]] = []
        self.removals: list[tuple[str, str]] = []

    def upsert_for_version(
        self, spec_id: str, version: str, embeddings: list[list[float]]
    ) -> None:
        self.upserts.append((spec_id, version, embeddings))

    def remove_for_version(self, spec_id: str, version: str) -> None:
        self.removals.append((spec_id, version))


def _word_bytes() -> bytes:
    document = Document()
    document.add_heading("2 Scope", level=1)
    document.add_paragraph("Word section content.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _presentation_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, text in enumerate(
        (
            "Contents",
            "1 Scope",
            "Table 1: Equipment",
            "slideonlysignal",
            _LONG_SLIDE_TEXT,
        )
    ):
        shape = slide.shapes.add_textbox(
            Inches(1), Inches(index + 1), Inches(8), Inches(0.5)
        )
        shape.text_frame.text = text

    table_shape = slide.shapes.add_table(
        2, 2, Inches(1), Inches(6), Inches(8), Inches(1)
    )
    table = table_shape.table
    table.cell(0, 0).text = "Equipment"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "NR-ARFCN"
    table.cell(1, 1).text = "supported"
    slide.notes_slide.notes_text_frame.text = "speaker secret notesonlyprobe"

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_speaker_note_search_phrase_occurs_only_in_notes() -> None:
    slide = Presentation(io.BytesIO(_presentation_bytes())).slides[0]
    slide_text = " ".join(
        shape.text_frame.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    notes_text = slide.notes_slide.notes_text_frame.text

    assert "speaker secret" not in slide_text
    assert "speaker secret" in notes_text


def _zip_bytes(entries: list[tuple[str, bytes]]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, payload in entries:
            archive.writestr(name, payload)
    return buffer.getvalue()


def _mixed_zip() -> bytes:
    return _zip_bytes(
        [
            ("slides.pptx", _presentation_bytes()),
            ("38331-j00.docx", _word_bytes()),
        ]
    )


def _word_only_zip() -> bytes:
    return _zip_bytes([("38331-j00.docx", _word_bytes())])


def _settings(
    cache_dir,
    *,
    chunk_size: int = 8,
    chunk_overlap: int = 0,
    max_chunk_chars: int = 80,
    auto_embed_on_parse: bool = False,
):
    settings = get_settings()
    return settings.model_copy(
        update={
            "semantic_search": settings.semantic_search.model_copy(
                update={
                    "chunk_size": chunk_size,
                    "chunk_overlap": chunk_overlap,
                }
            ),
            "spec_doc": settings.spec_doc.model_copy(
                update={
                    "cache_dir": cache_dir,
                    "max_chunk_chars": max_chunk_chars,
                    "chunk_overlap": chunk_overlap,
                    "auto_index_on_parse": True,
                    "auto_embed_on_parse": auto_embed_on_parse,
                }
            ),
        }
    )


def _service(
    tmp_path,
    payload: bytes,
    *,
    doc_repo=None,
    inventory_repo=None,
    search_service=None,
    semantic_service=None,
    settings=None,
    fetcher=None,
) -> SpecDocService:
    return SpecDocService(
        spec_repo=_SpecRepo(),
        doc_repo=doc_repo,
        inventory_repo=inventory_repo,
        search_service=search_service,
        semantic_service=semantic_service,
        settings=settings or _settings(tmp_path / "settings-cache"),
        fetcher=fetcher or (lambda _url: payload),
        cache_dir=tmp_path / "zip-cache",
    )


def test_mixed_zip_persists_full_presentation_and_indexes_its_text(
    sqlite_env, tmp_path
):
    create_schema("all")
    doc_repo = SQLAlchemySpecDocRepository()
    search = SpecDocSearchService()
    settings = _settings(
        tmp_path / "settings-cache",
        chunk_size=8,
        chunk_overlap=0,
        max_chunk_chars=80,
        auto_embed_on_parse=True,
    )
    embedder = _RecordingEmbedder()
    vector_repo = _RecordingVectorRepository()
    semantic = SpecDocSemanticService(
        fts5_service=search,
        embedder=embedder,
        vector_repo=vector_repo,
        settings=settings,
        doc_repo=doc_repo,
    )
    service = _service(
        tmp_path,
        _mixed_zip(),
        doc_repo=doc_repo,
        search_service=search,
        semantic_service=semantic,
        settings=settings,
    )

    source = service.parse(_SPEC_ID)
    toc = service.get_toc(_SPEC_ID, _VERSION)
    chunks = service.list_chunks(_SPEC_ID, version=_VERSION)
    presentation_chunks = [
        chunk for chunk in chunks if chunk.source_file == "slides.pptx"
    ]
    presentation_chunk_ids = {chunk.chunk_id for chunk in presentation_chunks}
    presentations = service.list_presentations(_SPEC_ID, _VERSION)
    presentation = service.get_presentation(_SPEC_ID, _VERSION, "slides.pptx")
    filters = SpecDocSearchFilters(spec_id=_SPEC_ID, version=_VERSION)
    note_hits = search.search("notesonlyprobe", filters)
    slide_hits = search.search("slideonlysignal", filters)

    assert source.docx_count == 2  # Historical field name counts Word and PPTX files.
    assert presentation.markdown == _EXPECTED_MARKDOWN
    assert presentations == [presentation]
    assert presentation.file_order == 1
    assert presentation_chunks
    assert all(
        chunk.sections is None and chunk.tables is None
        for chunk in presentation_chunks
    )
    assert all(
        chunk.file_order == presentation.file_order for chunk in presentation_chunks
    )
    assert all(len(chunk.text.split()) <= 8 for chunk in presentation_chunks)
    assert all(len(chunk.text) <= 80 for chunk in presentation_chunks)
    assert all(entry.source_file != "slides.pptx" for entry in toc.entries)
    assert all(file.source_file != "slides.pptx" for file in toc.files)
    assert any(
        hit.chunk_id in presentation_chunk_ids and "notesonlyprobe" in hit.text
        for hit in note_hits
    )
    assert any(
        hit.chunk_id in presentation_chunk_ids and "slideonlysignal" in hit.text
        for hit in slide_hits
    )
    assert any("notesonlyprobe" in text for text in embedder.texts)
    assert any("slideonlysignal" in text for text in embedder.texts)
    assert vector_repo.upserts
    assert len(vector_repo.upserts[0][2]) == len(chunks)

    with pytest.raises(SpecDocError) as exc_info:
        service.get_presentation(_SPEC_ID, _VERSION, "missing.pptx")
    assert _SPEC_ID in str(exc_info.value)
    assert _VERSION in str(exc_info.value)
    assert "missing.pptx" in str(exc_info.value)


def test_full_presentation_detail_is_not_reconstructed_from_overlapping_chunks(
    sqlite_env, tmp_path
):
    create_schema("all")
    settings = _settings(
        tmp_path / "overlap-settings-cache",
        chunk_size=8,
        chunk_overlap=2,
        max_chunk_chars=256,
    )
    service = _service(tmp_path, _mixed_zip(), settings=settings)

    service.parse(_SPEC_ID)
    presentation = service.get_presentation(_SPEC_ID, _VERSION, "slides.pptx")
    chunks = [
        chunk
        for chunk in service.list_chunks(_SPEC_ID, version=_VERSION)
        if chunk.source_file == "slides.pptx"
    ]

    assert presentation.markdown == _EXPECTED_MARKDOWN
    assert len(chunks) > 1
    assert any(
        previous.text.split()[-2:] == following.text.split()[:2]
        for previous, following in pairwise(chunks)
    )
    assert len(presentation.markdown) > max(len(chunk.text) for chunk in chunks)


def test_zero_slide_presentation_persists_empty_markdown_without_a_chunk(
    sqlite_env, tmp_path
):
    create_schema("all")
    raw = _zip_bytes([("empty.pptx", build_empty_presentation())])
    service = _service(tmp_path, raw)

    source = service.parse(_SPEC_ID)
    presentation = service.get_presentation(_SPEC_ID, _VERSION, "empty.pptx")

    assert source.docx_count == 1
    assert source.chunk_count == 0
    assert presentation.markdown == ""
    assert service.list_presentations(_SPEC_ID, _VERSION) == [presentation]
    assert service.list_chunks(_SPEC_ID, version=_VERSION) == []


@pytest.mark.filterwarnings("ignore:Duplicate name:UserWarning")
def test_duplicate_presentation_paths_fail_but_inventory_keeps_both_entries(
    sqlite_env, tmp_path
):
    create_schema("all")
    inventory_repo = SQLAlchemySpecDocZipInventoryRepository()
    raw = _zip_bytes(
        [("slides.pptx", _presentation_bytes()), ("slides.pptx", b"second copy")]
    )
    service = _service(tmp_path, raw, inventory_repo=inventory_repo)

    with pytest.raises(ValueError, match=r"slides\.pptx"):
        service.parse(_SPEC_ID)

    inventory = inventory_repo.get(_SPEC_ID, _VERSION)
    assert inventory.zip_files == ("slides.pptx", "slides.pptx")


def test_bad_presentation_does_not_replace_previously_stored_version_content(
    sqlite_env, tmp_path
):
    create_schema("all")
    doc_repo = SQLAlchemySpecDocRepository()
    inventory_repo = SQLAlchemySpecDocZipInventoryRepository()
    service = _service(
        tmp_path,
        _mixed_zip(),
        doc_repo=doc_repo,
        inventory_repo=inventory_repo,
    )
    service.parse(_SPEC_ID)
    old_chunks = doc_repo.list_chunks(_SPEC_ID, version=_VERSION, limit=100)
    old_presentations = doc_repo.list_presentations(_SPEC_ID, _VERSION)
    old_toc = doc_repo.get_toc(_SPEC_ID, _VERSION)
    service._fetcher = lambda _url: _zip_bytes(
        [("broken.pptx", b"not a presentation")]
    )

    with pytest.raises(RuntimeError, match="presentation"):
        service.parse(_SPEC_ID, force=True)

    assert doc_repo.list_chunks(_SPEC_ID, version=_VERSION, limit=100) == old_chunks
    assert doc_repo.list_presentations(_SPEC_ID, _VERSION) == old_presentations
    assert doc_repo.get_toc(_SPEC_ID, _VERSION) == old_toc


def test_forced_word_only_reparse_removes_old_presentation_records(sqlite_env, tmp_path):
    create_schema("all")
    service = _service(tmp_path, _mixed_zip())
    service.parse(_SPEC_ID)
    assert service.list_presentations(_SPEC_ID, _VERSION)

    service._fetcher = lambda _url: _word_only_zip()
    source = service.parse(_SPEC_ID, force=True)

    assert source.docx_count == 1
    assert service.list_presentations(_SPEC_ID, _VERSION) == []
    assert all(
        chunk.source_file != "slides.pptx"
        for chunk in service.list_chunks(_SPEC_ID, version=_VERSION)
    )
    with pytest.raises(SpecDocError) as exc_info:
        service.get_presentation(_SPEC_ID, _VERSION, "slides.pptx")
    assert _SPEC_ID in str(exc_info.value)
    assert _VERSION in str(exc_info.value)
    assert "slides.pptx" in str(exc_info.value)


def test_parsed_version_without_force_does_not_download_or_backfill_presentations(
    sqlite_env, tmp_path
):
    create_schema("all")
    doc_repo = SQLAlchemySpecDocRepository()
    doc_repo.record_download(
        _SPEC_ID,
        _VERSION,
        release="Rel-18",
        ftp_url=_FTP_URL,
        docx_count=1,
    )
    doc_repo.record_parsed(_SPEC_ID, _VERSION, chunk_count=0)
    downloads: list[str] = []
    service = _service(
        tmp_path,
        _mixed_zip(),
        doc_repo=doc_repo,
        fetcher=lambda url: downloads.append(url) or _mixed_zip(),
    )

    source = service.parse(_SPEC_ID)

    assert source.parsed_at is not None
    assert downloads == []
    assert service.list_presentations(_SPEC_ID, _VERSION) == []


@pytest.mark.semantic
def test_mixed_presentation_slide_and_notes_use_real_sqlite_vec(sqlite_env, tmp_path):
    import numpy as np

    from doc3gpp.storage.repositories.spec_doc_vector_sql import (
        SQLAlchemySpecDocVectorRepository,
    )

    create_schema("all")
    doc_repo = SQLAlchemySpecDocRepository()
    search = SpecDocSearchService()
    settings = _settings(
        tmp_path / "semantic-settings-cache",
        chunk_size=8,
        chunk_overlap=0,
        max_chunk_chars=80,
    )
    vector_repo = SQLAlchemySpecDocVectorRepository()
    model_name = "presentation-acceptance"
    vector_repo.reset_for_rebuild(3, model_name)

    class _Embedder:
        dim = 3
        model_name = "presentation-acceptance"

        def encode(self, texts: list[str]) -> np.ndarray:
            vectors = np.zeros((len(texts), self.dim), dtype=np.float32)
            for row, value in enumerate(texts):
                lowered = value.casefold()
                slide_text = "slideonlysignal" in lowered
                notes_text = "notesonlyprobe" in lowered
                if slide_text:
                    vectors[row, 0] = 1.0
                if notes_text:
                    vectors[row, 1] = 1.0
                if not slide_text and not notes_text:
                    vectors[row, 2] = 1.0
            return vectors

    semantic = SpecDocSemanticService(
        fts5_service=search,
        embedder=_Embedder(),
        vector_repo=vector_repo,
        settings=settings,
        doc_repo=doc_repo,
    )
    service = _service(
        tmp_path,
        _mixed_zip(),
        doc_repo=doc_repo,
        search_service=search,
        semantic_service=semantic,
        settings=settings,
    )

    source = service.parse(_SPEC_ID)
    assert source.docx_count == 2
    presentation = service.get_presentation(_SPEC_ID, _VERSION, "slides.pptx")
    assert presentation.markdown == _EXPECTED_MARKDOWN
    presentation_chunk_ids = {
        chunk.chunk_id
        for chunk in service.list_chunks(_SPEC_ID, version=_VERSION, limit=100)
        if chunk.source_file == "slides.pptx"
    }
    assert presentation_chunk_ids
    semantic.index_for_version(_SPEC_ID, _VERSION)

    filters = SpecDocSearchFilters(spec_id=_SPEC_ID, version=_VERSION)
    for query in ("slideonlysignal", "notesonlyprobe"):
        fts_hits = search.search(query, filters)
        assert any(
            hit.chunk_id in presentation_chunk_ids and query in hit.text.casefold()
            for hit in fts_hits
        )

        vector_hits = semantic.search(
            query,
            fts5_query=None,
            filters=filters,
            limit=10,
            fts5_weight=0.5,
        )
        assert vector_hits
        assert vector_hits[0].chunk_id in presentation_chunk_ids
        assert vector_hits[0].rank_vec == 0
        assert vector_hits[0].hit is not None
        assert query in vector_hits[0].hit.text.casefold()


def test_force_reparse_refreshes_presentation_and_preserves_word_content(
    sqlite_env, tmp_path, monkeypatch,
):
    from typer.testing import CliRunner

    from doc3gpp.cli import app

    create_schema("all")
    doc_repo = SQLAlchemySpecDocRepository()
    service = _service(
        tmp_path,
        _mixed_zip(),
        doc_repo=doc_repo,
    )
    service.parse(_SPEC_ID)
    old_presentation = service.get_presentation(
        _SPEC_ID, _VERSION, "slides.pptx"
    )
    old_word_chunks = [
        chunk
        for chunk in doc_repo.list_chunks(_SPEC_ID, version=_VERSION, limit=100)
        if chunk.source_file == "38331-j00.docx"
    ]
    old_toc = doc_repo.get_toc(_SPEC_ID, _VERSION)

    updated = Presentation(io.BytesIO(_presentation_bytes()))
    slide = updated.slides[0]
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text == "slideonlysignal":
            shape.text_frame.text = "refreshed slide acceptance signal"
            break
    slide.notes_slide.notes_text_frame.text = "refreshed speaker notes"
    updated_buffer = io.BytesIO()
    updated.save(updated_buffer)
    refreshed_zip = _zip_bytes(
        [
            ("slides.pptx", updated_buffer.getvalue()),
            ("38331-j00.docx", _word_bytes()),
        ]
    )
    service._fetcher = lambda _url: refreshed_zip
    monkeypatch.setattr("doc3gpp.cli.build_spec_doc_service", lambda **_kwargs: service)

    result = CliRunner().invoke(
        app,
        ["spec", "doc", "parse", "--spec", _SPEC_ID, "--force"],
    )

    assert result.exit_code == 0, result.output
    source = service.get_source(_SPEC_ID, _VERSION)
    assert source is not None
    assert source.docx_count == 2
    new_presentation = service.get_presentation(
        _SPEC_ID, _VERSION, "slides.pptx"
    )
    assert new_presentation.markdown != old_presentation.markdown
    assert "refreshed slide acceptance signal" in new_presentation.markdown
    assert "refreshed speaker notes" in new_presentation.markdown
    new_word_chunks = [
        chunk
        for chunk in doc_repo.list_chunks(_SPEC_ID, version=_VERSION, limit=100)
        if chunk.source_file == "38331-j00.docx"
    ]
    assert new_word_chunks == old_word_chunks
    new_toc = doc_repo.get_toc(_SPEC_ID, _VERSION)
    assert new_toc.entries == old_toc.entries
    assert new_toc.files == old_toc.files
    assert new_toc.docx_count == old_toc.docx_count


def test_legacy_ppt_member_uses_impress_or_deterministic_converter(
    sqlite_env, tmp_path, monkeypatch,
):
    import shutil
    import subprocess

    from doc3gpp.parsers import presentation_converter

    create_schema("all")
    doc_repo = SQLAlchemySpecDocRepository()
    pptx_bytes = _presentation_bytes()
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        monkeypatch.setattr(
            presentation_converter,
            "_as_pptx_bytes",
            lambda _payload, _filename: pptx_bytes,
        )
        legacy_ppt = b"deterministic legacy PPT fixture"
    else:
        source_dir = tmp_path / "pptx-source"
        output_dir = tmp_path / "legacy-ppt"
        source_dir.mkdir()
        output_dir.mkdir()
        source = source_dir / "presentation.pptx"
        source.write_bytes(pptx_bytes)
        result = subprocess.run(
            [
                executable,
                f"-env:UserInstallation={(tmp_path / 'legacy-export-profile').as_uri()}",
                "--headless",
                "--nologo",
                "--nodefault",
                "--norestore",
                "--convert-to",
                "ppt",
                "--outdir",
                str(output_dir),
                str(source),
            ],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
        assert result.returncode == 0, result.stderr or result.stdout
        legacy_path = output_dir / "presentation.ppt"
        assert legacy_path.is_file(), result.stderr or result.stdout
        legacy_ppt = legacy_path.read_bytes()

    raw = _zip_bytes(
        [
            ("slides/legacy.ppt", legacy_ppt),
            ("38331-j00.docx", _word_bytes()),
        ]
    )
    service = _service(
        tmp_path,
        raw,
        doc_repo=doc_repo,
        settings=_settings(tmp_path / "legacy-settings-cache"),
    )

    source = service.parse(_SPEC_ID)
    converted = service.get_presentation(_SPEC_ID, _VERSION, "slides/legacy.ppt")

    assert source.docx_count == 2
    assert converted.source_file == "slides/legacy.ppt"
    assert converted.markdown.startswith("## Slide 1")
    assert "Contents" in converted.markdown
    assert "Table 1: Equipment" in converted.markdown

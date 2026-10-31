"""Optional real binary-Word round trip; ordinary tests double only LibreOffice."""

from __future__ import annotations

import shutil
import subprocess

import pytest

from doc3gpp.parsers.docx_converter import (
    HeadingBlock,
    TableBlock,
    convert_document_to_blocks,
    convert_document_to_markdown,
)
from doc3gpp.parsers.presentation_converter import convert_presentation_to_markdown

OFFICE = shutil.which("libreoffice") or shutil.which("soffice")


@pytest.mark.skipif(OFFICE is None, reason="optional LibreOffice Writer executable is not installed")
def test_real_legacy_word_preserves_structured_content(tmp_path):
    pytest.importorskip("docx")
    from tests.fixtures.docx_builders import build_golden_docx

    source = tmp_path / "example.docx"
    source.write_bytes(build_golden_docx())
    completed = subprocess.run(
        [
            OFFICE, f"-env:UserInstallation={(tmp_path / 'export-profile').as_uri()}",
            "--headless", "--convert-to", "doc:MS Word 97", "--outdir", str(tmp_path), str(source),
        ],
        capture_output=True, text=True, timeout=120, check=True,
    )
    legacy = tmp_path / "example.doc"
    assert legacy.is_file(), completed.stdout + completed.stderr
    payload = legacy.read_bytes()
    assert payload.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))  # Actual OLE .doc, not renamed OOXML.

    markdown = convert_document_to_markdown(payload, "example.doc")
    assert "1 Scope" in markdown
    assert "Spanned body" in markdown
    assert "inserted text" in markdown
    assert "deleted text" in markdown
    blocks = convert_document_to_blocks(payload, "example.doc")
    assert any(isinstance(block, HeadingBlock) and block.section_no == "1" for block in blocks)
    assert any(isinstance(block, TableBlock) and "Header A" in block.gfm for block in blocks)



@pytest.mark.skipif(OFFICE is None, reason="optional LibreOffice Impress executable is not installed")
def test_real_legacy_presentation_roundtrip_extracts_slide_text(tmp_path):
    pytest.importorskip("pptx")
    import io

    from pptx import Presentation
    from pptx.util import Inches

    from tests.fixtures.presentation_builders import build_presentation_with_content

    def export_to_ppt(source, output_dir, profile):
        return subprocess.run(
            [
                OFFICE,
                f"-env:UserInstallation={profile.as_uri()}",
                "--headless",
                "--convert-to",
                "ppt:MS PowerPoint 97",
                "--outdir",
                str(output_dir),
                str(source),
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )

    probe_dir = tmp_path / "impress-probe"
    probe_output = probe_dir / "out"
    probe_dir.mkdir()
    probe_output.mkdir()
    probe_presentation = Presentation()
    probe_slide = probe_presentation.slides.add_slide(probe_presentation.slide_layouts[6])
    probe_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = (
        "Impress conversion probe"
    )
    probe_source = probe_dir / "probe.pptx"
    probe_buffer = io.BytesIO()
    probe_presentation.save(probe_buffer)
    probe_source.write_bytes(probe_buffer.getvalue())
    probe_result = export_to_ppt(
        probe_source, probe_output, probe_dir / "profile"
    )
    if probe_result.returncode != 0:
        if "source file could not be loaded" in probe_result.stderr.lower():
            pytest.skip("LibreOffice Impress import filters are not installed")
        pytest.fail(
            f"LibreOffice Impress probe failed: {probe_result.stdout}{probe_result.stderr}"
        )
    assert (probe_output / "probe.ppt").is_file(), (
        probe_result.stdout + probe_result.stderr
    )

    source = tmp_path / "slides.pptx"
    source.write_bytes(build_presentation_with_content())
    completed = export_to_ppt(source, tmp_path, tmp_path / "export-profile")
    assert completed.returncode == 0, completed.stdout + completed.stderr
    legacy = tmp_path / "slides.ppt"
    assert legacy.is_file(), completed.stdout + completed.stderr
    payload = legacy.read_bytes()
    assert payload.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))

    markdown = convert_presentation_to_markdown(payload, "slides/slides.ppt")
    assert "## Slide 1" in markdown
    assert "日本語" in markdown
    assert "group text" in markdown
    assert "second slide text" in markdown

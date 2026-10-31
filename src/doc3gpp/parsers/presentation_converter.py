"""Convert PowerPoint presentations to deterministic Markdown.

``python-pptx`` and LibreOffice are optional. The Python package is imported
only when a presentation is converted, and LibreOffice is used only for legacy
``.ppt`` input.
"""

from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
from collections.abc import Iterable
from pathlib import Path

from doc3gpp.parsers.word_converter import _stop_conversion
from doc3gpp.settings.loader import get_settings


class PythonPptxNotInstalledError(ImportError):
    """Raised when presentation extraction is requested without ``python-pptx``."""

    def __init__(self, message: str | None = None) -> None:
        super().__init__(
            message
            or (
                "python-pptx is required for presentation extraction. "
                "Install with `pip install doc3gpp[extract]`."
            )
        )


def _presentation_library():
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise PythonPptxNotInstalledError() from exc
    return Presentation, MSO_SHAPE_TYPE


def _as_pptx_bytes(payload: bytes, filename: str) -> bytes:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        raise RuntimeError(
            f"LibreOffice Impress is required to parse legacy .ppt source {filename!r}. "
            "Install LibreOffice Impress and put 'libreoffice' or 'soffice' on PATH. "
            ".pptx parsing does not require LibreOffice."
        )

    timeout_seconds = get_settings().word_conversion.timeout_seconds
    with tempfile.TemporaryDirectory(prefix="doc3gpp-ppt-") as temporary:
        root = Path(temporary)
        source = root / "input.ppt"
        source.write_bytes(payload)
        output = root / "input.pptx"
        command = [
            executable,
            f"-env:UserInstallation={(root / 'profile').as_uri()}",
            "--headless",
            "--nologo",
            "--nodefault",
            "--norestore",
            "--convert-to",
            "pptx:Impress MS PowerPoint 2007 XML",
            "--outdir",
            str(root),
            str(source),
        ]
        try:
            process = subprocess.Popen(
                command,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                encoding="utf-8",
                errors="replace",
                start_new_session=os.name == "posix",
            )
            try:
                stdout, stderr = process.communicate(timeout=timeout_seconds)
            except BaseException:
                _stop_conversion(process)
                raise
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError(
                f"LibreOffice conversion timed out after {timeout_seconds}s for "
                f"{filename!r}"
            ) from exc
        except OSError as exc:
            raise RuntimeError(f"Cannot run LibreOffice for {filename!r}: {exc}") from exc

        if process.returncode != 0 or not output.is_file() or output.stat().st_size == 0:
            detail = (stderr or stdout or "no PPTX output").strip()[:500]
            raise RuntimeError(
                f"LibreOffice failed to convert {filename!r} "
                f"(exit {process.returncode}): {detail}"
            )
        return output.read_bytes()


def _cell_markdown(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    return "<br>".join(normalized.split("\n")).replace("|", r"\|")


def _table_markdown(table) -> str:
    rows = [
        "| " + " | ".join(_cell_markdown(cell.text) for cell in row.cells) + " |"
        for row in table.rows
    ]
    if not rows:
        return ""
    separator = "| " + " | ".join("---" for _ in table.columns) + " |"
    return "\n".join((rows[0], separator, *rows[1:]))


def _shape_markdown(shapes: Iterable, shape_types) -> list[str]:
    content: list[str] = []
    for shape in shapes:
        if shape.shape_type == shape_types.GROUP:
            content.extend(_shape_markdown(shape.shapes, shape_types))
        elif getattr(shape, "has_table", False):
            table = _table_markdown(shape.table)
            if table:
                content.append(table)
        elif getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.replace("\v", "\n").strip()
            if text:
                content.append(text)
    return content


def convert_presentation_to_markdown(payload: bytes, filename: str) -> str:
    """Convert ``.pptx`` or legacy ``.ppt`` bytes to ordered Markdown.

    Slide text, tables, recursively grouped shapes, and notes already present in
    the source are included. Images, charts, and other graphical content are
    intentionally ignored. A valid zero-slide presentation returns ``""``.
    """
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pptx", ".ppt"}:
        raise ValueError(f"Unsupported extension {suffix!r}; expected .pptx or .ppt")

    pptx_bytes = _as_pptx_bytes(payload, filename) if suffix == ".ppt" else payload
    Presentation, shape_types = _presentation_library()
    try:
        presentation = Presentation(io.BytesIO(pptx_bytes))
    except Exception as exc:
        raise RuntimeError(
            f"python-pptx could not read presentation {filename!r}: {exc}"
        ) from exc

    slides: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        content = [f"## Slide {slide_number}"]
        content.extend(_shape_markdown(slide.shapes, shape_types))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                content.append(f"**Speaker notes:**\n\n{notes}")
        slides.append("\n\n".join(content))
    return "\n\n".join(slides)


__all__ = ["PythonPptxNotInstalledError", "convert_presentation_to_markdown"]

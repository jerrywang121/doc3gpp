# tests/unit/test_spec_doc_blocks.py
import io
from docx import Document
from doc3gpp.parsers.docx_converter import convert_document_to_blocks, HeadingBlock, ParagraphBlock, TableBlock

def _make_docx():
    doc = Document()
    doc.add_heading("1 Scope", level=1)
    doc.add_paragraph("Some text here.")
    doc.add_paragraph("Table 1: My caption")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(0, 1).text = "B"
    t.cell(1, 0).text = "C"
    t.cell(1, 1).text = "D"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _convert_docx_with_paragraphs(paragraphs):
    doc = Document()
    for style, text in paragraphs:
        doc.add_paragraph(text, style=style)
    buf = io.BytesIO()
    doc.save(buf)
    return convert_document_to_blocks(buf.getvalue(), "in-memory.docx")


def _convert_docx_with_table_caption(caption):
    doc = Document()
    doc.add_paragraph(caption)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Value"
    buf = io.BytesIO()
    doc.save(buf)
    return convert_document_to_blocks(buf.getvalue(), "in-memory.docx")

def test_blocks_heading_paragraph_table():
    blocks = convert_document_to_blocks(_make_docx(), "38331-j30.docx")
    assert isinstance(blocks[0], HeadingBlock) and blocks[0].level == 1
    assert blocks[0].section_no == "1" and blocks[0].title == "Scope"
    assert any(isinstance(b, ParagraphBlock) for b in blocks)
    tables = [b for b in blocks if isinstance(b, TableBlock)]
    assert len(tables) == 1 and tables[0].table_no == "1"


def test_heading_accepts_alphanumeric_section_suffixes():
    blocks = _convert_docx_with_paragraphs(
        [
            ("Heading 1", "7.2A.3 Scope"),
            ("Heading 2", "7.2A.3A Extended details"),
        ]
    )
    headings = [block for block in blocks if isinstance(block, HeadingBlock)]
    assert [(h.section_no, h.title) for h in headings] == [
        ("7.2A.3", "Scope"),
        ("7.2A.3A", "Extended details"),
    ]


def test_table_caption_accepts_3gpp_identifier_and_separator():
    blocks = _convert_docx_with_table_caption("Table 7.2A.3A: UE capability values")
    table = next(block for block in blocks if isinstance(block, TableBlock))
    assert table.table_no == "7.2A.3A"
    assert table.table_title == "UE capability values"

def test_rejects_dot_doc():
    import pytest
    with pytest.raises(ValueError):
        convert_document_to_blocks(b"xx", "old.doc")

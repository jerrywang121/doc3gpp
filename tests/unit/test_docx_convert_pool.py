"""Tests for the bounded spawn-context DOCX conversion pool.

``workers <= 1`` runs in-process and must match :class:`InlineDocxPool`.
``workers >= 2`` runs real spawned worker processes; those tests stay few so
the suite remains fast. Broken-pool behaviour is driven through fake
executors, so no worker process is killed.
"""

from __future__ import annotations

import io
import multiprocessing
from concurrent.futures import Future
from concurrent.futures.process import BrokenProcessPool

import pytest
from docx import Document

from doc3gpp.parsers.docx_converter import (
    InlineDocxPool,
    convert_document_to_markdown,
)
from doc3gpp.services.docx_convert_pool import DocxConvertPool, DocxConvertPoolError
from tests.fixtures.docx_builders import build_golden_branches_docx, build_golden_docx


def _golden_with_extra_paragraphs(count: int) -> bytes:
    document = Document(io.BytesIO(build_golden_docx()))
    for index in range(count):
        document.add_paragraph(f"Extra paragraph {index}.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _empty_docx() -> bytes:
    buffer = io.BytesIO()
    Document().save(buffer)
    return buffer.getvalue()


def _three_members() -> list[tuple[str, bytes]]:
    return [
        ("golden.docx", build_golden_docx()),
        ("branches.docx", build_golden_branches_docx()),
        ("golden_extra.docx", _golden_with_extra_paragraphs(5)),
    ]


def _five_sized_members() -> list[tuple[str, bytes]]:
    return [
        ("golden.docx", build_golden_docx()),
        ("golden_5.docx", _golden_with_extra_paragraphs(5)),
        ("branches.docx", build_golden_branches_docx()),
        ("golden_50.docx", _golden_with_extra_paragraphs(50)),
        ("golden_500.docx", _golden_with_extra_paragraphs(500)),
    ]


@pytest.fixture(autouse=True)
def _no_worker_processes_left_behind():
    yield
    assert multiprocessing.active_children() == []


@pytest.mark.parametrize("workers", [0, 1])
def test_in_process_pool_matches_inline_pool_and_creates_no_executor(workers, monkeypatch) -> None:
    def no_executor(self):
        raise AssertionError("workers <= 1 must not create an executor")

    monkeypatch.setattr(DocxConvertPool, "_get_executor", no_executor)
    members = _three_members()
    pool = DocxConvertPool(workers)
    assert pool.convert_members("markdown", members) == InlineDocxPool().convert_members(
        "markdown", members
    )
    assert pool.convert_members("blocks", members) == InlineDocxPool().convert_members(
        "blocks", members
    )


def test_two_workers_match_inline_pool_for_sized_members_in_input_order() -> None:
    members = _five_sized_members()
    expected_markdown = InlineDocxPool().convert_members("markdown", members)
    expected_blocks = InlineDocxPool().convert_members("blocks", members)
    with DocxConvertPool(2) as pool:
        assert pool.convert_members("markdown", members) == expected_markdown
        assert pool.convert_members("blocks", members) == expected_blocks


def test_bad_member_is_an_exception_outcome_and_the_call_does_not_raise() -> None:
    golden = build_golden_docx()
    with DocxConvertPool(2) as pool:
        outcomes = pool.convert_members(
            "markdown",
            [("a.docx", golden), ("bad.docx", b"junk"), ("c.docx", golden)],
        )
    assert isinstance(outcomes[1], Exception)
    assert outcomes[0] == convert_document_to_markdown(golden, "a.docx")
    assert outcomes[2] == convert_document_to_markdown(golden, "c.docx")


def test_worker_exceptions_round_trip_with_same_type_and_message() -> None:
    members = [("legacy.doc", build_golden_docx()), ("empty.docx", _empty_docx())]
    expected = InlineDocxPool().convert_members("markdown", members)
    with DocxConvertPool(2) as pool:
        outcomes = pool.convert_members("markdown", members)
    assert type(outcomes[0]) is ValueError
    assert type(outcomes[1]) is RuntimeError
    for got, want in zip(outcomes, expected, strict=True):
        assert type(got) is type(want)
        assert str(got) == str(want)


def test_negative_workers_is_rejected() -> None:
    with pytest.raises(ValueError):
        DocxConvertPool(-1)


def test_close_twice_does_not_raise_and_leaves_no_worker_processes() -> None:
    pool = DocxConvertPool(2)
    pool.convert_members("markdown", [("a.docx", build_golden_docx())])
    pool.close()
    pool.close()


class _CountingFuture(Future):
    def __init__(self, owner: _FakeExecutor) -> None:
        super().__init__()
        self._owner = owner

    def result(self, timeout=None):
        self._owner.outstanding -= 1
        return super().result(timeout)


class _FakeExecutor:
    def __init__(self, *, broken_on_submit: bool = False, broken_on_result: bool = False) -> None:
        self.broken_on_submit = broken_on_submit
        self.broken_on_result = broken_on_result
        self.shutdown_calls = 0
        self.outstanding = 0
        self.max_outstanding = 0

    @property
    def shut_down(self) -> bool:
        return self.shutdown_calls > 0

    def submit(self, fn, *args):
        if self.broken_on_submit:
            raise BrokenProcessPool("worker terminated abruptly")
        self.outstanding += 1
        self.max_outstanding = max(self.max_outstanding, self.outstanding)
        future = _CountingFuture(self)
        if self.broken_on_result:
            future.set_exception(BrokenProcessPool("worker terminated abruptly"))
        else:
            future.set_result(fn(*args))
        return future

    def shutdown(self, wait: bool = True, **_kwargs) -> None:
        self.shutdown_calls += 1


def _use_fake_executors(monkeypatch, *executors: _FakeExecutor) -> None:
    pending = list(executors)
    original = DocxConvertPool._get_executor

    def get_executor(self):
        if self._executor is None and pending:
            self._executor = pending.pop(0)
        return original(self)

    monkeypatch.setattr(DocxConvertPool, "_get_executor", get_executor)


def test_broken_submit_fails_member_shuts_executor_and_next_call_recreates(monkeypatch) -> None:
    golden = build_golden_docx()
    broken = _FakeExecutor(broken_on_submit=True)
    fresh = _FakeExecutor()
    _use_fake_executors(monkeypatch, broken, fresh)
    with DocxConvertPool(2) as pool:
        first = pool.convert_members("markdown", [("bad_pool.docx", golden)])
        assert isinstance(first[0], DocxConvertPoolError)
        assert str(first[0]) == (
            "DOCX conversion pool broke while converting bad_pool.docx: "
            "worker terminated abruptly"
        )
        assert broken.shut_down

        second = pool.convert_members("markdown", [("ok.docx", golden)])
    assert second == [convert_document_to_markdown(golden, "ok.docx")]
    assert fresh.shut_down


def test_broken_submit_spares_later_members_of_the_same_call(monkeypatch) -> None:
    golden = build_golden_docx()
    broken = _FakeExecutor(broken_on_submit=True)
    fresh = _FakeExecutor()
    _use_fake_executors(monkeypatch, broken, fresh)
    with DocxConvertPool(2) as pool:
        outcomes = pool.convert_members(
            "markdown", [("bad_pool.docx", golden), ("ok.docx", golden)]
        )
    assert isinstance(outcomes[0], DocxConvertPoolError)
    assert outcomes[1] == convert_document_to_markdown(golden, "ok.docx")
    assert broken.shut_down


def test_discarded_executor_is_shut_down_exactly_once(monkeypatch) -> None:
    golden = build_golden_docx()
    broken = _FakeExecutor(broken_on_result=True)
    _use_fake_executors(monkeypatch, broken)
    with DocxConvertPool(2) as pool:
        outcomes = pool.convert_members("markdown", [("a.docx", golden), ("b.docx", golden)])
        assert all(isinstance(outcome, DocxConvertPoolError) for outcome in outcomes)
        assert broken.shutdown_calls == 1
    assert broken.shutdown_calls == 1


def test_submit_time_break_shuts_executor_down_exactly_once(monkeypatch) -> None:
    golden = build_golden_docx()
    broken = _FakeExecutor(broken_on_submit=True)
    fresh = _FakeExecutor()
    _use_fake_executors(monkeypatch, broken, fresh)
    with DocxConvertPool(2) as pool:
        outcomes = pool.convert_members(
            "markdown", [("a.docx", golden), ("b.docx", golden), ("c.docx", golden)]
        )
    assert isinstance(outcomes[0], DocxConvertPoolError)
    assert outcomes[1:] == [
        convert_document_to_markdown(golden, "b.docx"),
        convert_document_to_markdown(golden, "c.docx"),
    ]
    assert broken.shutdown_calls == 1
    assert fresh.shutdown_calls == 1


def test_window_keeps_two_per_worker_outstanding_and_preserves_input_order(monkeypatch) -> None:
    members = [(f"m{index}.docx", _golden_with_extra_paragraphs(index)) for index in range(8)]
    fake = _FakeExecutor()
    _use_fake_executors(monkeypatch, fake)
    with DocxConvertPool(2) as pool:
        outcomes = pool.convert_members("markdown", members)
    assert fake.max_outstanding == 4
    assert outcomes == [convert_document_to_markdown(payload, name) for name, payload in members]


def test_broken_pool_fails_only_the_outstanding_window(monkeypatch) -> None:
    members = [(f"m{index}.docx", _golden_with_extra_paragraphs(index)) for index in range(10)]
    broken = _FakeExecutor(broken_on_result=True)
    fresh = _FakeExecutor()
    _use_fake_executors(monkeypatch, broken, fresh)
    with DocxConvertPool(2) as pool:
        outcomes = pool.convert_members("markdown", members)
    assert all(isinstance(outcome, DocxConvertPoolError) for outcome in outcomes[:4])
    assert outcomes[4:] == [
        convert_document_to_markdown(payload, name) for name, payload in members[4:]
    ]
    assert broken.shutdown_calls == 1


def test_broken_result_fails_member_shuts_executor_and_next_call_recreates(monkeypatch) -> None:
    golden = build_golden_docx()
    broken = _FakeExecutor(broken_on_result=True)
    fresh = _FakeExecutor()
    _use_fake_executors(monkeypatch, broken, fresh)
    with DocxConvertPool(2) as pool:
        first = pool.convert_members("markdown", [("bad_pool.docx", golden)])
        assert isinstance(first[0], DocxConvertPoolError)
        assert broken.shut_down

        second = pool.convert_members("markdown", [("ok.docx", golden)])
    assert second == [convert_document_to_markdown(golden, "ok.docx")]
    assert fresh.shut_down

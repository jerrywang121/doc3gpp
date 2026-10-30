"""Integration tests for SpecDocService fetch/parse/toc (Task 9)."""
import io
import zipfile
from unittest.mock import MagicMock

import pytest
from docx import Document

from doc3gpp.models.spec import SpecVersion
from doc3gpp.models.spec_doc import (
    SpecDocNoDocxError,
    SpecDocUnknownVersionError,
)
from doc3gpp.models.spec_doc_zip_inventory import SpecDocZipInventory
from doc3gpp.services.spec_doc_service import SpecDocService
from doc3gpp.storage.db.migrate import create_schema
from doc3gpp.storage.repositories.spec_doc_zip_inventory_sql import (
    SQLAlchemySpecDocZipInventoryRepository,
)
from tests.fixtures.presentation_builders import build_presentation_with_content


class FakeSpecRepo:
    def __init__(self, versions):
        self._v = versions

    def list_versions(self, spec_id, **kw):
        return self._v.get(spec_id, [])


def _make_zip_bytes(paragraphs=("hello handover world",), name="38331-j30.docx"):
    doc = Document()
    doc.add_heading("1 Scope", level=1)
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    docx = buf.getvalue()
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w") as z:
        z.writestr(name, docx)
    return zb.getvalue()


def _make_inventory_zip_bytes():
    with zipfile.ZipFile(io.BytesIO(_make_zip_bytes())) as source:
        docx = source.read("38331-j30.docx")
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w") as archive:
        archive.writestr("main.docx", docx)
        archive.writestr("directory/", b"")
        archive.writestr("readme.txt", b"readme")
        archive.writestr(".hidden", b"hidden")
        archive.writestr("__MACOSX/meta", b"metadata")
        archive.writestr("nested.zip", b"not opened as an archive")
        archive.writestr("日本語/文書.docx", docx)
        archive.writestr(r"folder\backslash.txt", b"separator preserved")
        archive.writestr("readme.txt", b"duplicate")
    return zb.getvalue()


def _make_chunk_context_zip_bytes():
    doc = Document()
    doc.add_heading("1 Scope", level=1)
    doc.add_paragraph("Body text for the scope.")
    doc.add_paragraph("Table 1: Values")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "A"
    table.cell(1, 0).text = "value"

    buf = io.BytesIO()
    doc.save(buf)
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w") as z:
        z.writestr("38331-j30.docx", buf.getvalue())
    return zb.getvalue()


def _versions():
    return {
        "38.331": [
            SpecVersion(
                "38.331",
                "18.5.0",
                "https://www.3gpp.org/ftp/x/38331.zip",
                "Rel-18",
            )
        ]
    }


def test_parse_many_buckets_and_identity_skip(sqlite_env):
    zip_bytes = _make_zip_bytes()
    versions = _versions()
    svc = SpecDocService(
        spec_repo=FakeSpecRepo(versions),
        fetcher=lambda url: zip_bytes,
        cache_dir=sqlite_env.parent / "speccache",
    )
    create_schema("all")
    r1 = svc.parse_many(["38.331"])
    assert "38.331" in r1.successes
    r2 = svc.parse_many(["38.331"])
    assert "38.331" in r2.skipped  # immutable: second parse skips
    toc = svc.get_toc("38.331", "18.5.0")
    assert toc.entries and toc.entries[0].section_no == "1"


@pytest.mark.parametrize("force", [False, True])
def test_reparse_includes_doc_members_and_refreshes_document_count(
    sqlite_env, tmp_path, monkeypatch, force,
):
    from doc3gpp.services.docx_convert_pool import DocxConvertPool
    from tests.fixtures.legacy_doc import install_fake_libreoffice

    create_schema("all")
    original = _make_zip_bytes()
    with zipfile.ZipFile(io.BytesIO(original)) as archive:
        docx = archive.read("38331-j30.docx")
    install_fake_libreoffice(monkeypatch, tmp_path, {b"legacy spec": docx})
    mixed = io.BytesIO()
    with zipfile.ZipFile(mixed, "w") as archive:
        archive.writestr("legacy/38331-j30.DOC", b"legacy spec")
        archive.writestr("38331-j30.docx", docx)

    with DocxConvertPool(2) as pool:
        service = SpecDocService(
            spec_repo=FakeSpecRepo(_versions()), fetcher=lambda _url: original,
            cache_dir=tmp_path / "cache", docx_pool=pool,
        )
        assert service.parse("38.331").docx_count == 1
        service._fetcher = lambda _url: mixed.getvalue()
        if not force:
            service._doc_repo.invalidate_parse_state("38.331", "18.5.0")
            service._write_zip_cache("38.331", "18.5.0", mixed.getvalue())
        source = service.parse("38.331", force=force)

    assert source.docx_count == 2  # Historical field name now counts both Word formats.
    assert source.parsed_at is not None
    toc = service.get_toc("38.331", "18.5.0")
    assert toc.docx_count == 2
    assert {entry.source_file for entry in toc.entries} == {"legacy/38331-j30.DOC", "38331-j30.docx"}
    assert all(entry.section_no == "1" for entry in toc.entries)
    chunks = service.list_chunks("38.331", version="18.5.0")
    assert len(chunks) == 2
    assert all("hello handover world" in chunk.text for chunk in chunks)
    inventory = SQLAlchemySpecDocZipInventoryRepository().get("38.331", "18.5.0")
    assert inventory.zip_files == ("legacy/38331-j30.DOC", "38331-j30.docx")


def test_parse_counts_presentations_without_toc_or_section_metadata(sqlite_env):
    create_schema("all")
    original = _make_zip_bytes()
    with zipfile.ZipFile(io.BytesIO(original)) as archive:
        docx = archive.read("38331-j30.docx")

    mixed = io.BytesIO()
    with zipfile.ZipFile(mixed, "w") as archive:
        archive.writestr("slides.pptx", build_presentation_with_content())
        archive.writestr("38331-j30.docx", docx)

    service = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        fetcher=lambda _url: mixed.getvalue(),
        cache_dir=sqlite_env.parent / "presentation-count-cache",
    )

    source = service.parse("38.331")
    toc = service.get_toc("38.331", "18.5.0")
    presentation_chunks = [
        chunk
        for chunk in service.list_chunks("38.331", version="18.5.0")
        if chunk.source_file == "slides.pptx"
    ]

    assert source.docx_count == 2  # Historical field name counts Word and presentations.
    assert toc.docx_count == 2
    assert all(entry.source_file != "slides.pptx" for entry in toc.entries)
    assert all(file.source_file != "slides.pptx" for file in toc.files)
    assert presentation_chunks
    assert all(
        chunk.sections is None and chunk.tables is None
        for chunk in presentation_chunks
    )


def test_parse_stores_chunk_content_and_context(sqlite_env):
    create_schema("all")
    svc = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        fetcher=lambda url: _make_chunk_context_zip_bytes(),
        cache_dir=sqlite_env.parent / "speccache",
    )

    svc.parse("38.331")
    chunks = svc.list_chunks("38.331", version="18.5.0")

    assert chunks
    chunk = chunks[0]
    assert "# 1 Scope" in chunk.text
    assert "Table 1: Values" in chunk.text
    assert "| A |" in chunk.text
    assert chunk.sections == "1 Scope"
    assert chunk.tables == "1 Values"


def test_post_chunk_cache_failure_does_not_mark_source_parsed(sqlite_env, monkeypatch):
    create_schema("all")
    svc = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        fetcher=lambda url: _make_zip_bytes(),
        cache_dir=sqlite_env.parent / "speccache",
    )
    original_write = svc._write_markdown_cache
    calls = 0

    def fail_once(spec_id, version, ordered):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise OSError("markdown cache unavailable")
        return original_write(spec_id, version, ordered)

    monkeypatch.setattr(svc, "_write_markdown_cache", fail_once)

    first = svc.parse_many(["38.331"])
    assert "38.331" in first.failures
    source = svc.get_source("38.331", "18.5.0")
    assert source is not None and source.parsed_at is None

    second = svc.parse_many(["38.331"])
    assert "38.331" in second.successes
    assert calls == 2


def test_failed_force_reparse_invalidates_old_marker_for_retry(sqlite_env, monkeypatch):
    create_schema("all")
    svc = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        fetcher=lambda url: _make_zip_bytes(),
        cache_dir=sqlite_env.parent / "speccache",
    )
    first = svc.parse_many(["38.331"])
    assert "38.331" in first.successes
    source = svc.get_source("38.331", "18.5.0")
    assert source is not None and source.parsed_at is not None

    original_write = svc._write_markdown_cache
    calls = 0

    def fail_once(spec_id, version, ordered):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise OSError("markdown cache unavailable")
        return original_write(spec_id, version, ordered)

    monkeypatch.setattr(svc, "_write_markdown_cache", fail_once)

    forced = svc.parse_many(["38.331"], force=True)
    assert "38.331" in forced.failures
    source = svc.get_source("38.331", "18.5.0")
    assert source is not None and source.parsed_at is None

    retry = svc.parse_many(["38.331"])
    assert "38.331" in retry.successes
    assert calls == 2


def test_failed_force_download_invalidates_old_marker_for_retry(sqlite_env, monkeypatch):
    create_schema("all")
    svc = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        fetcher=lambda url: _make_zip_bytes(),
        cache_dir=sqlite_env.parent / "speccache",
    )
    first = svc.parse_many(["38.331"])
    assert "38.331" in first.successes

    def fail_fetch(url):
        raise OSError("upstream unavailable")

    monkeypatch.setattr(svc, "_fetcher", fail_fetch)
    forced = svc.parse_many(["38.331"], force=True)
    assert "38.331" in forced.failures
    source = svc.get_source("38.331", "18.5.0")
    assert source is not None and source.parsed_at is None

    monkeypatch.setattr(svc, "_fetcher", lambda url: _make_zip_bytes())
    retry = svc.parse_many(["38.331"])
    assert "38.331" in retry.successes


def test_fetch_skip_uses_zip_cache_without_network(sqlite_env):
    create_schema("all")
    zip_bytes = _make_zip_bytes()
    cache_dir = sqlite_env.parent / "speccache"
    svc = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        fetcher=lambda url: zip_bytes,
        cache_dir=cache_dir,
    )
    svc.parse_many(["38.331"])

    def _boom(url):
        raise AssertionError(f"network must not be touched: {url}")

    svc2 = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()), fetcher=_boom, cache_dir=cache_dir
    )
    src = svc2.fetch("38.331")
    assert src.version == "18.5.0" and src.parsed_at is not None


def test_oversized_zip_lands_in_skipped(sqlite_env):
    from doc3gpp.settings.schema import Settings, SpecDocSettings

    create_schema("all")
    settings = Settings(spec_doc=SpecDocSettings(max_zip_size_kb=1))
    svc = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        fetcher=lambda url: _make_zip_bytes(),
        cache_dir=sqlite_env.parent / "speccache",
        settings=settings,
    )
    result = svc.parse_many(["38.331"])
    assert "38.331" in result.skipped
    assert "38.331" not in result.successes


def test_unknown_spec_lands_in_failures(sqlite_env):
    create_schema("all")
    svc = SpecDocService(
        spec_repo=FakeSpecRepo({}),
        fetcher=lambda url: b"",
        cache_dir=sqlite_env.parent / "speccache",
    )
    result = svc.parse_many(["99.999"])
    assert "99.999" in result.failures
    assert "spec sync" in result.failures["99.999"]


def test_unknown_version_lands_in_failures(sqlite_env):
    create_schema("all")
    svc = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        fetcher=lambda url: _make_zip_bytes(),
        cache_dir=sqlite_env.parent / "speccache",
    )
    result = svc.parse_many(["38.331"], version="99.0")
    assert "38.331" in result.failures


def test_no_docx_lands_in_failures(sqlite_env):
    create_schema("all")
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w") as z:
        z.writestr("readme.txt", b"no docx here")
    svc = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        fetcher=lambda url: zb.getvalue(),
        cache_dir=sqlite_env.parent / "speccache",
    )
    result = svc.parse_many(["38.331"])
    assert "38.331" in result.failures
    assert "no .docx" in result.failures["38.331"]
    with pytest.raises(SpecDocNoDocxError):
        svc.parse("38.331", force=True)


def test_force_reparses_and_redownloads(sqlite_env):
    create_schema("all")
    calls: list[str] = []

    def _counting_fetcher(url):
        calls.append(url)
        return _make_zip_bytes()

    svc = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        fetcher=_counting_fetcher,
        cache_dir=sqlite_env.parent / "speccache",
    )
    r1 = svc.parse_many(["38.331"])
    assert "38.331" in r1.successes
    assert len(calls) == 1
    r2 = svc.parse_many(["38.331"], force=True)
    assert "38.331" in r2.successes
    assert len(calls) == 2


def test_get_toc_miss_raises(sqlite_env):
    create_schema("all")
    svc = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        fetcher=lambda url: _make_zip_bytes(),
        cache_dir=sqlite_env.parent / "speccache",
    )
    with pytest.raises(SpecDocUnknownVersionError):
        svc.get_toc("38.331", "18.5.0")


@pytest.mark.filterwarnings("ignore:Duplicate name:UserWarning")
def test_prepare_carries_inventory_without_writing(sqlite_env):
    create_schema("all")
    inventory_repo = MagicMock()
    service = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        inventory_repo=inventory_repo,
        fetcher=lambda _url: _make_inventory_zip_bytes(),
        cache_dir=sqlite_env.parent / "inventory-prepare-cache",
    )

    prepared = service.prepare_parse("38.331")

    assert prepared.inventory == SpecDocZipInventory(
        "38.331",
        "18.5.0",
        "Rel-18",
        (
            "main.docx",
            "readme.txt",
            ".hidden",
            "__MACOSX/meta",
            "nested.zip",
            "日本語/文書.docx",
            r"folder\backslash.txt",
            "readme.txt",
        ),
    )
    inventory_repo.upsert.assert_not_called()


def test_no_docx_and_empty_zip_publish_inventory_but_fail(sqlite_env):
    create_schema("all")
    no_docx = io.BytesIO()
    with zipfile.ZipFile(no_docx, "w") as archive:
        archive.writestr("readme.txt", b"no docx here")
    empty = io.BytesIO()
    with zipfile.ZipFile(empty, "w"):
        pass

    inventory_repo = SQLAlchemySpecDocZipInventoryRepository()
    for index, (raw, expected_names) in enumerate(
        ((no_docx.getvalue(), ("readme.txt",)), (empty.getvalue(), ()))
    ):
        service = SpecDocService(
            spec_repo=FakeSpecRepo(_versions()),
            inventory_repo=inventory_repo,
            fetcher=lambda _url, payload=raw: payload,
            cache_dir=sqlite_env.parent / f"inventory-empty-cache-{index}",
        )
        result = service.parse_many(["38.331"])

        assert "38.331" in result.failures
        assert "no .docx" in result.failures["38.331"]
        assert inventory_repo.get("38.331", "18.5.0").zip_files == expected_names
        assert service.get_source("38.331", "18.5.0") is None


def test_immutable_skip_does_not_download_enumerate_or_write(
    sqlite_env, monkeypatch
):
    from doc3gpp.services import spec_doc_service as service_module

    create_schema("all")
    inventory_repo = SQLAlchemySpecDocZipInventoryRepository()
    writes = []
    original_upsert = inventory_repo.upsert

    def record_upsert(inventory):
        writes.append(inventory)
        original_upsert(inventory)

    monkeypatch.setattr(inventory_repo, "upsert", record_upsert)
    service = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        inventory_repo=inventory_repo,
        fetcher=lambda _url: _make_zip_bytes(),
        cache_dir=sqlite_env.parent / "immutable-inventory-cache",
    )
    assert "38.331" in service.parse_many(["38.331"]).successes
    assert len(writes) == 1

    def fail_fetch(_url):
        raise AssertionError("immutable skip must not fetch")

    def fail_enumeration(_raw):
        raise AssertionError("immutable skip must not enumerate")

    service._fetcher = fail_fetch
    monkeypatch.setattr(service_module, "enumerate_zip_files", fail_enumeration)

    result = service.parse_many(["38.331"])

    assert "38.331" in result.skipped
    assert len(writes) == 1


def test_force_reparse_replaces_inventory(sqlite_env):
    create_schema("all")
    payloads = [
        _make_zip_bytes(name="first.docx"),
        _make_zip_bytes(name="second.docx"),
    ]
    calls = []

    def fetcher(url):
        calls.append(url)
        return payloads[len(calls) - 1]

    inventory_repo = SQLAlchemySpecDocZipInventoryRepository()
    service = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        inventory_repo=inventory_repo,
        fetcher=fetcher,
        cache_dir=sqlite_env.parent / "force-inventory-cache",
    )

    assert "38.331" in service.parse_many(["38.331"]).successes
    assert inventory_repo.get("38.331", "18.5.0").zip_files == ("first.docx",)
    assert "38.331" in service.parse_many(["38.331"], force=True).successes

    assert len(calls) == 2
    assert inventory_repo.get("38.331", "18.5.0").zip_files == ("second.docx",)


def test_fetch_helper_alone_does_not_collect_inventory(sqlite_env):
    create_schema("all")
    inventory_repo = SQLAlchemySpecDocZipInventoryRepository()
    service = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        inventory_repo=inventory_repo,
        fetcher=lambda _url: _make_zip_bytes(),
        cache_dir=sqlite_env.parent / "fetch-inventory-cache",
    )

    service.fetch("38.331")

    assert inventory_repo.get("38.331", "18.5.0") is None


def test_unparsed_cached_zip_without_fetch_persists_inventory(sqlite_env):
    create_schema("all")
    inventory_repo = SQLAlchemySpecDocZipInventoryRepository()
    service = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        inventory_repo=inventory_repo,
        fetcher=lambda _url: (_ for _ in ()).throw(
            AssertionError("cached unparsed ZIP must not fetch")
        ),
        cache_dir=sqlite_env.parent / "cached-inventory-cache",
    )
    service._write_zip_cache("38.331", "18.5.0", _make_zip_bytes())

    result = service.parse_many(["38.331"])

    assert "38.331" in result.successes
    assert inventory_repo.get("38.331", "18.5.0").zip_files == ("38331-j30.docx",)


@pytest.mark.parametrize("failure", ("download", "oversized", "invalid"))
def test_pre_parse_failure_leaves_previous_inventory_unchanged(
    sqlite_env, failure
):
    from doc3gpp.settings.schema import Settings, SpecDocSettings

    create_schema("all")
    inventory_repo = SQLAlchemySpecDocZipInventoryRepository()
    previous = SpecDocZipInventory(
        "38.331", "18.5.0", "Rel-18", ("previous.docx",)
    )
    inventory_repo.upsert(previous)

    def fetcher(_url):
        if failure == "download":
            raise OSError("upstream unavailable")
        if failure == "oversized":
            return _make_zip_bytes()
        return b"not a ZIP"

    settings = (
        Settings(spec_doc=SpecDocSettings(max_zip_size_kb=1))
        if failure == "oversized"
        else Settings()
    )
    service = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        inventory_repo=inventory_repo,
        settings=settings,
        fetcher=fetcher,
        cache_dir=sqlite_env.parent / f"pre-parse-error-{failure}",
    )

    result = service.parse_many(["38.331"])

    assert ("38.331" in result.skipped) is (failure == "oversized")
    assert ("38.331" in result.failures) is (failure != "oversized")
    assert inventory_repo.get("38.331", "18.5.0") == previous
    assert service.get_source("38.331", "18.5.0") is None


def test_synchronous_forced_conversion_failure_invalidates_old_marker(sqlite_env):
    create_schema("all")
    inventory_repo = SQLAlchemySpecDocZipInventoryRepository()
    service = SpecDocService(
        spec_repo=FakeSpecRepo(_versions()),
        inventory_repo=inventory_repo,
        fetcher=lambda _url: _make_zip_bytes(),
        cache_dir=sqlite_env.parent / "forced-failure-cache",
    )
    assert "38.331" in service.parse_many(["38.331"]).successes
    error = RuntimeError("conversion exploded")

    class FailedPool:
        def convert_members(self, _kind, members):
            return [error for _ in members]

    service._docx_pool = FailedPool()
    result = service.parse_many(["38.331"], force=True)

    assert result.failures["38.331"] == "conversion exploded"
    source = service.get_source("38.331", "18.5.0")
    assert source is not None and source.parsed_at is None
    assert inventory_repo.get("38.331", "18.5.0").zip_files == ("38331-j30.docx",)

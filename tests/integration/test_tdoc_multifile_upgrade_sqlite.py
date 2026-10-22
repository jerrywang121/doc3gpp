"""SQLite integration coverage for legacy and incomplete TDoc extractions."""

from __future__ import annotations

import hashlib
import io
import json
import zipfile
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from sqlalchemy import delete, text, update

from doc3gpp.models.search import SearchFilters
from doc3gpp.models.tdoc import TDoc
from doc3gpp.models.tdoc_cr import TDocCRDetails, TDocExtractMeta
from doc3gpp.models.tdoc_document import TDocDocument
from doc3gpp.scraping.cache import TDocCache
from doc3gpp.services.search_service import SearchService
from doc3gpp.services.tdoc_cr_service import TDocCrService
from doc3gpp.storage.db.migrate import create_schema
from doc3gpp.storage.db.models import TDocDocumentOrm, TDocExtractOrm
from doc3gpp.storage.db.session import get_engine, get_session_factory
from doc3gpp.storage.repositories.search_sql import SQLAlchemySearchIndexRepository
from doc3gpp.storage.repositories.tdoc_cr_change_details_sql import (
    SQLAlchemyTDocCrChangeDetailsRepository,
)
from doc3gpp.storage.repositories.tdoc_cr_sql import SQLAlchemyTDocCrRepository
from doc3gpp.storage.repositories.tdoc_cr_ttcn_sql import (
    SQLAlchemyTDocCrTtcnRepository,
)
from doc3gpp.storage.repositories.tdoc_extraction_sql import (
    SQLAlchemyTDocExtractionRepository,
)
from doc3gpp.storage.repositories.tdoc_sql import SQLAlchemyTDocRepository

FIXTURES_DIR = Path(__file__).resolve().parents[1] / "fixtures" / "tdoc_cr_doc"


def _document(
    ftp_url: str,
    source_file: str,
    file_order: int,
    markdown: str,
) -> TDocDocument:
    document_id = hashlib.sha256(
        json.dumps(
            [ftp_url, source_file],
            ensure_ascii=False,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    markdown_file = (
        f"documents/{document_id}/{source_file.rsplit('/', 1)[-1].rsplit('.', 1)[0]}.md"
    )
    return TDocDocument(
        document_id=document_id,
        source_file=source_file,
        markdown_file=markdown_file,
        file_order=file_order,
        markdown_text=markdown,
        content_hash=hashlib.sha256(markdown.encode("utf-8")).hexdigest(),
    )


def _seed_tdoc(
    tdoc_id: str,
    ftp_url: str,
    *,
    meeting_id: int | None = None,
) -> None:
    SQLAlchemyTDocRepository().upsert(
        TDoc(
            tdoc_id=tdoc_id,
            type="CR",
            ftp_url=ftp_url,
            meeting_id=meeting_id,
        )
    )


def _seed_v2_bundle(
    tdoc_id: str,
    ftp_url: str,
    sources: tuple[tuple[str, str], ...],
) -> tuple[TDocDocument, ...]:
    details = TDocCRDetails(
        tdoc_id=tdoc_id,
        ftp_url=ftp_url,
        title=f"CR {tdoc_id}",
    )
    documents = tuple(
        _document(ftp_url, source_file, order, markdown)
        for order, (source_file, markdown) in enumerate(sources)
    )
    meta = TDocExtractMeta(
        ftp_url=ftp_url,
        tdoc_id=tdoc_id,
        cache_file=f"{tdoc_id}.zip",
        doc_filename=sources[0][0],
        extraction_version=2,
        document_count=len(documents),
    )
    SQLAlchemyTDocExtractionRepository().replace_bundle(
        details,
        None,
        None,
        meta,
        documents,
    )
    return documents


def _docx_available() -> bool:
    try:
        from docx import Document  # noqa: F401
    except ImportError:
        return False
    return True


def _zip_with_support_docs(
    main_zip: bytes,
    support_texts: tuple[str, ...],
) -> bytes:
    from docx import Document

    with zipfile.ZipFile(io.BytesIO(main_zip)) as archive:
        main_name = next(
            name for name in archive.namelist() if name.lower().endswith(".docx")
        )
        main_docx = archive.read(main_name)

    def support_docx(text: str) -> bytes:
        document = Document()
        document.add_paragraph(text)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("z/R5s260009.docx", main_docx)
        for index, text in enumerate(support_texts):
            archive.writestr(
                f"{'a' if index == 0 else 'b'}/shared.docx",
                support_docx(text),
            )
    return output.getvalue()


def test_pending_query_requeues_legacy_and_incomplete_v2_bundles(sqlite_env) -> None:
    create_schema()
    tdoc_repo = SQLAlchemyTDocRepository()
    cr_repo = SQLAlchemyTDocCrRepository()

    legacy_id = "R5s260101"
    legacy_url = "docs/r5s260101.zip"
    _seed_tdoc(legacy_id, legacy_url)
    cr_repo.upsert(
        TDocCRDetails(tdoc_id=legacy_id, ftp_url=legacy_url, title="Legacy CR")
    )
    cr_repo.upsert_extract_meta(
        TDocExtractMeta(
            ftp_url=legacy_url,
            tdoc_id=legacy_id,
            cache_file=f"{legacy_id}.zip",
            doc_filename="z/main.docx",
            extraction_version=1,
            document_count=0,
        )
    )

    complete_id = "R5s260102"
    complete_url = "docs/r5s260102.zip"
    _seed_tdoc(complete_id, complete_url)
    _seed_v2_bundle(
        complete_id,
        complete_url,
        (("z/main.docx", "complete main"), ("a/support.docx", "complete support")),
    )

    missing_child_id = "R5s260103"
    missing_child_url = "docs/r5s260103.zip"
    _seed_tdoc(missing_child_id, missing_child_url)
    _seed_v2_bundle(
        missing_child_id,
        missing_child_url,
        (("z/main.docx", "main"), ("a/support.docx", "support")),
    )

    missing_main_id = "R5s260104"
    missing_main_url = "docs/r5s260104.zip"
    _seed_tdoc(missing_main_id, missing_main_url)
    _seed_v2_bundle(
        missing_main_id,
        missing_main_url,
        (("z/main.docx", "main"), ("a/support.docx", "support")),
    )

    session_factory = get_session_factory()
    with session_factory() as session:
        session.execute(
            delete(TDocDocumentOrm).where(
                TDocDocumentOrm.ftp_url == missing_child_url,
                TDocDocumentOrm.source_file == "a/support.docx",
            )
        )
        session.execute(
            update(TDocExtractOrm)
            .where(TDocExtractOrm.ftp_url == missing_main_url)
            .values(doc_filename="missing/main.docx")
        )
        session.commit()

    pending_ids = {
        tdoc.tdoc_id
        for tdoc in tdoc_repo.list(limit=20, offset=0, exclude_parsed=True)
    }

    assert pending_ids == {legacy_id, missing_child_id, missing_main_id}


@pytest.mark.skipif(
    not _docx_available(),
    reason="python-docx not installed; install with `pip install doc3gpp[extract]`",
)
def test_forced_reparse_removes_support_rows_missing_from_new_archive(
    sqlite_env,
    tmp_path: Path,
) -> None:
    create_schema()
    fixture = FIXTURES_DIR / "R5s260009.zip"
    first_zip = _zip_with_support_docs(
        fixture.read_bytes(),
        ("old support alpha", "old support beta"),
    )
    second_zip = _zip_with_support_docs(fixture.read_bytes(), ())

    tdoc_id = "R5s260009"
    ftp_url = "docs/r5s260009.zip"
    source_url = "https://www.3gpp.org/ftp/Docs/R5s260009.zip"
    _seed_tdoc(tdoc_id, ftp_url)

    scraper = MagicMock()
    scraper.get_bytes.side_effect = [first_zip, second_zip]
    tdoc_repo = SQLAlchemyTDocRepository()
    service = TDocCrService(
        cache=TDocCache(root=tmp_path / "cache", size_limit_bytes=0),
        scraper_client=scraper,
        cr_repository=SQLAlchemyTDocCrRepository(),
        cr_ttcn_repository=SQLAlchemyTDocCrTtcnRepository(),
        cr_change_details_repository=SQLAlchemyTDocCrChangeDetailsRepository(),
        tdoc_repository=tdoc_repo,
        extraction_repository=SQLAlchemyTDocExtractionRepository(),
    )

    first = service.extract_from_url(source_url)
    assert [document.source_file for document in first.documents] == [
        "z/R5s260009.docx",
        "a/shared.docx",
        "b/shared.docx",
    ]

    second = service.extract_from_url(source_url, force=True)
    assert second.from_cache is False
    assert [document.source_file for document in second.documents] == [
        "z/R5s260009.docx",
    ]
    assert [
        document.source_file
        for document in SQLAlchemyTDocExtractionRepository().list_documents(ftp_url)
    ] == ["z/R5s260009.docx"]
    assert scraper.get_bytes.call_count == 2


def test_index_rebuild_reads_sql_content_after_cache_purge_without_download(
    sqlite_env,
    tmp_path: Path,
    monkeypatch,
) -> None:
    create_schema()
    with get_engine().begin() as connection:
        connection.execute(
            text(
                "INSERT OR IGNORE INTO tsgs (tsg_name, short_name, description) "
                "VALUES ('RAN WG1', 'RAN1', 'TSG RAN WG1')"
            )
        )
        connection.execute(
            text(
                """
                INSERT OR IGNORE INTO meetings (
                    meeting_id, name, title, location, tsg, start_date,
                    end_date, ftp_url, tdoc_list_last_sync
                ) VALUES (
                    100, 'RAN1#100', 'RAN1#100', 'Online', 'RAN1',
                    '2026-01-01', '2026-01-05',
                    'https://www.3gpp.org/ftp/meetings/RAN1_100',
                    '2026-01-05T00:00:00'
                )
                """
            )
        )

    tdoc_id = "R5s260105"
    ftp_url = "docs/r5s260105.zip"
    _seed_tdoc(tdoc_id, ftp_url, meeting_id=100)
    _seed_v2_bundle(
        tdoc_id,
        ftp_url,
        (
            ("z/main.docx", "main CR text"),
            ("support/annex.docx", "sqlcontentaftercachepurge"),
        ),
    )

    cache = TDocCache(root=tmp_path / "cache", size_limit_bytes=0)
    cache.put_bytes("R5s260105.zip", b"zip cache", "zips")
    cache.put_bytes("R5s260105.zip", b"markdown cache", "markdown")
    assert cache.purge_subdir("zips") == 1
    assert cache.purge_subdir("markdown") == 1

    def unexpected_download(*_args, **_kwargs):
        pytest.fail("an index rebuild must not download a TDoc archive")

    monkeypatch.setattr(
        "doc3gpp.services.tdoc_cr_service.download_tdoc_zip",
        unexpected_download,
    )
    search_repo = SQLAlchemySearchIndexRepository()
    search_service = SearchService(repo=search_repo, reranker=object())
    list(
        search_service.rebuild(
            batch_size=10,
            resume=False,
            stale_only=False,
            quiet=True,
        )
    )

    hits = search_repo.search("sqlcontentaftercachepurge", SearchFilters(limit=10))
    assert len(hits) == 1
    assert hits[0].matched_documents[0].source_file == "support/annex.docx"



def test_pending_selection_scopes_completeness_to_current_ftp_url(sqlite_env) -> None:
    create_schema()
    tdoc_repo = SQLAlchemyTDocRepository()
    tdoc_id = "R5s260106"
    old_url = "docs/old/r5s260106.zip"
    current_url = "docs/current/r5s260106.zip"
    sources = (("z/main.docx", "main body"), ("support/annex.docx", "support body"))

    _seed_tdoc(tdoc_id, old_url)
    _seed_v2_bundle(tdoc_id, old_url, sources)
    _seed_tdoc(tdoc_id, current_url)

    pending_before_current_extract = {
        tdoc.tdoc_id
        for tdoc in tdoc_repo.list(limit=20, offset=0, exclude_parsed=True)
    }
    assert pending_before_current_extract == {tdoc_id}

    _seed_v2_bundle(tdoc_id, current_url, sources)
    pending_after_current_extract = {
        tdoc.tdoc_id
        for tdoc in tdoc_repo.list(limit=20, offset=0, exclude_parsed=True)
    }
    assert pending_after_current_extract == set()

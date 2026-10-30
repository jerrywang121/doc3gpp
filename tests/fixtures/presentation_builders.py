"""Generated PowerPoint fixtures for presentation conversion tests."""

from __future__ import annotations

import io

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def _presentation_bytes(presentation) -> bytes:
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def build_empty_presentation() -> bytes:
    """Return a valid presentation with no slides."""
    return _presentation_bytes(Presentation())


def build_presentation_with_content() -> bytes:
    """Return two slides covering text, nested groups, a table, and notes."""
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    first_text = first.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    )
    first_text.text_frame.text = "日本語"

    group = first.shapes.add_group_shape()
    nested_group = group.shapes.add_group_shape()
    grouped_text = nested_group.shapes.add_textbox(
        Inches(1), Inches(2), Inches(4), Inches(1)
    )
    grouped_text.text_frame.text = "group text"

    table_shape = first.shapes.add_table(
        2, 1, Inches(1), Inches(3), Inches(5), Inches(1)
    )
    table_shape.table.cell(0, 0).text = "Value"
    table_shape.table.cell(1, 0).text = "a|b\nline1\nline2"

    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second_text = second.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    )
    second_text.text_frame.text = "second slide text"
    second.notes_slide.notes_text_frame.text = "speaker secret"
    return _presentation_bytes(presentation)


def build_presentation_without_notes() -> bytes:
    """Return one slide with text and no notes slide."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    text.text_frame.text = "slide without notes"
    return _presentation_bytes(presentation)


def build_presentation_with_graphics() -> bytes:
    """Return chart/image shapes whose graphical content must not be inferred."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    text.text_frame.text = "visible text"

    chart_data = CategoryChartData()
    chart_data.categories = ["category secret"]
    chart_data.add_series("series secret", (7,))
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(2),
        Inches(4),
        Inches(3),
        chart_data,
    )
    chart_shape.chart.has_title = True
    chart_shape.chart.chart_title.text_frame.text = "chart secret"

    image = Image.new("RGB", (140, 32), "white")
    ImageDraw.Draw(image).text((2, 2), "image secret", fill="black")
    image_bytes = io.BytesIO()
    image.save(image_bytes, format="PNG")
    slide.shapes.add_picture(
        io.BytesIO(image_bytes.getvalue()), Inches(6), Inches(2)
    )
    return _presentation_bytes(presentation)

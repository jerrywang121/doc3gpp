"""Unit tests for shared TDoc ZIP conversion with tolerant annexes."""

from __future__ import annotations

import hashlib
import io
import logging
import zipfile
from collections.abc import Callable, Sequence
from pathlib import Path

import pytest

from doc3gpp.models.tdoc_document import markdown_member_name
from doc3gpp.models.tdoc_identity import document_id_for
from doc3gpp.parsers.docx_converter import DocxConvertPoolError, convert_document_to_markdown
from doc3gpp.parsers.tdoc_zip import TDocZipLimitError, TDocZipValidationError
from doc3gpp.parsers.tdoc_zip_convert import (
    TDocZipConversion,
    convert_direct_main_markdown,
    convert_direct_payload,
    convert_tdoc_zip,
)

IDENTITY = "tsg_ran/wg5_test_ex-t1/tsgr5_100/docs/r5-123456.zip"


def _docx_available() -> bool:
    try:
        from docx import Document  # noqa: F401
    except ImportError:
        return False
    return True


def _save_docx(document) -> bytes:
    payload = io.BytesIO()
    document.save(payload)
    return payload.getvalue()


def _make_docx(text: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    return _save_docx(document)


def _make_empty_docx() -> bytes:
    from docx import Document

    return _save_docx(Document())


def _make_zip(files: list[tuple[str, bytes]]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, payload in files:
            archive.writestr(name, payload)
    return output.getvalue()


def _decode(payload: bytes, source_file: str) -> str:
    """Stub converter: the member payload is already the Markdown text."""
    return payload.decode("utf-8")


def _failing(*failing: str, message: str = "cannot convert") -> Callable[[bytes, str], str]:
    """Stub converter that raises ``message`` for the named sources and decodes the rest."""

    def convert(payload: bytes, source_file: str) -> str:
        if source_file in failing:
            raise RuntimeError(f"{message} {source_file}")
        return _decode(payload, source_file)

    return convert


def _convert(
    zip_bytes: bytes,
    *,
    convert: Callable[[bytes, str], str] = _decode,
    max_documents: int = 100,
) -> TDocZipConversion:
    return convert_tdoc_zip(
        zip_bytes,
        tdoc_id="R5-123456",
        identity=IDENTITY,
        max_documents=max_documents,
        max_uncompressed_bytes=0,
        convert=convert,
    )


@pytest.mark.skipif(
    not _docx_available(),
    reason="python-docx not installed; install with `pip install doc3gpp[extract]`",
)
def test_stores_main_then_annexes_in_order() -> None:
    zip_bytes = _make_zip(
        [
            ("annex/b.docx", _make_docx("annex b body")),
            ("R5-123456.docx", _make_docx("main body")),
            ("annex/a.docx", _make_docx("annex a body")),
        ]
    )

    result = _convert(zip_bytes, convert=convert_document_to_markdown)

    assert result.main_source_file == "R5-123456.docx"
    assert [document.source_file for document in result.documents] == [
        "R5-123456.docx",
        "annex/a.docx",
        "annex/b.docx",
    ]
    assert [document.file_order for document in result.documents] == [0, 1, 2]
    assert result.skipped_sources == ()
    for document in result.documents:
        assert document.document_id == document_id_for(IDENTITY, document.source_file)
        assert document.markdown_file == markdown_member_name(
            document.document_id, document.source_file
        )
        assert (
            document.content_hash
            == hashlib.sha256(document.markdown_text.encode("utf-8")).hexdigest()
        )
    assert "main body" in result.documents[0].markdown_text


def test_failing_annex_is_skipped_and_not_fatal(caplog: pytest.LogCaptureFixture) -> None:
    zip_bytes = _make_zip(
        [
            ("R5-123456.docx", b"main body"),
            ("annex/a.docx", b"annex a body"),
            ("annex/bad.docx", b"corrupt"),
            ("annex/c.docx", b"annex c body"),
        ]
    )

    with caplog.at_level(logging.WARNING):
        result = _convert(zip_bytes, convert=_failing("annex/bad.docx"))

    assert result.skipped_sources == ("annex/bad.docx",)
    assert [document.source_file for document in result.documents] == [
        "R5-123456.docx",
        "annex/a.docx",
        "annex/c.docx",
    ]
    assert [document.file_order for document in result.documents] == [0, 1, 2]
    assert any(
        record.levelno == logging.WARNING and "annex/bad.docx" in record.getMessage()
        for record in caplog.records
    )


def test_failing_main_raises() -> None:
    zip_bytes = _make_zip([("R5-123456.docx", b"main body"), ("annex/a.docx", b"annex")])

    with pytest.raises(RuntimeError, match="cannot convert R5-123456.docx"):
        _convert(zip_bytes, convert=_failing("R5-123456.docx"))


def test_doc_only_main_is_a_supported_conversion_source() -> None:
    zip_bytes = _make_zip([("R5-123456.doc", b"legacy binary")])

    result = _convert(zip_bytes)
    assert result.main_source_file == "R5-123456.doc"
    assert result.documents[0].markdown_text == "legacy binary"


def test_zero_text_documents_are_kept() -> None:
    zip_bytes = _make_zip([("R5-123456.docx", b""), ("annex/a.docx", b"")])

    result = _convert(zip_bytes, convert=lambda _payload, _source_file: "")

    assert [document.source_file for document in result.documents] == [
        "R5-123456.docx",
        "annex/a.docx",
    ]
    assert [document.markdown_text for document in result.documents] == ["", ""]
    empty_hash = hashlib.sha256(b"").hexdigest()
    assert [document.content_hash for document in result.documents] == [empty_hash, empty_hash]
    assert result.skipped_sources == ()


@pytest.mark.skipif(
    not _docx_available(),
    reason="python-docx not installed; install with `pip install doc3gpp[extract]`",
)
def test_empty_text_main_is_kept_as_empty_markdown() -> None:
    zip_bytes = _make_zip(
        [
            ("R5-123456.docx", _make_empty_docx()),
            ("annex/a.docx", _make_docx("annex a body")),
        ]
    )

    result = _convert(zip_bytes, convert=convert_document_to_markdown)

    assert result.main_source_file == "R5-123456.docx"
    assert [document.source_file for document in result.documents] == [
        "R5-123456.docx",
        "annex/a.docx",
    ]
    assert result.documents[0].markdown_text == ""
    assert result.documents[0].content_hash == hashlib.sha256(b"").hexdigest()
    assert result.skipped_sources == ()


@pytest.mark.skipif(
    not _docx_available(),
    reason="python-docx not installed; install with `pip install doc3gpp[extract]`",
)
def test_empty_text_annex_is_kept_not_skipped() -> None:
    zip_bytes = _make_zip(
        [
            ("R5-123456.docx", _make_docx("main body")),
            ("annex/a.docx", _make_empty_docx()),
        ]
    )

    result = _convert(zip_bytes, convert=convert_document_to_markdown)

    assert [document.source_file for document in result.documents] == [
        "R5-123456.docx",
        "annex/a.docx",
    ]
    assert result.documents[1].file_order == 1
    assert result.documents[1].markdown_text == ""
    assert result.documents[1].content_hash == hashlib.sha256(b"").hexdigest()
    assert result.skipped_sources == ()


def test_empty_text_error_message_is_kept_as_empty_markdown(
    caplog: pytest.LogCaptureFixture,
) -> None:
    zip_bytes = _make_zip(
        [
            ("R5-123456.docx", b"unused"),
            ("annex/a.docx", b"unused"),
            ("annex/b.docx", b"annex b body"),
        ]
    )
    convert = _failing(
        "R5-123456.docx",
        "annex/a.docx",
        message="python-docx returned empty result for",
    )

    with caplog.at_level(logging.WARNING):
        result = _convert(zip_bytes, convert=convert)

    assert [document.source_file for document in result.documents] == [
        "R5-123456.docx",
        "annex/a.docx",
        "annex/b.docx",
    ]
    assert [document.markdown_text for document in result.documents] == ["", "", "annex b body"]
    assert result.documents[0].content_hash == hashlib.sha256(b"").hexdigest()
    assert result.skipped_sources == ()
    assert [record for record in caplog.records if record.levelno >= logging.WARNING] == []


def test_other_runtime_error_on_annex_is_still_skipped(caplog: pytest.LogCaptureFixture) -> None:
    zip_bytes = _make_zip(
        [
            ("R5-123456.docx", b"main body"),
            ("annex/bad.docx", b"corrupt"),
            ("annex/c.docx", b"annex c body"),
        ]
    )

    with caplog.at_level(logging.WARNING):
        result = _convert(
            zip_bytes,
            convert=_failing("annex/bad.docx", message="python-docx failed on"),
        )

    assert result.skipped_sources == ("annex/bad.docx",)
    assert [document.source_file for document in result.documents] == [
        "R5-123456.docx",
        "annex/c.docx",
    ]
    assert any(
        record.levelno == logging.WARNING and "annex/bad.docx" in record.getMessage()
        for record in caplog.records
    )


def test_zip_document_limit_raises() -> None:
    zip_bytes = _make_zip([("R5-123456.docx", b"main"), ("annex/a.docx", b"annex")])

    with pytest.raises(TDocZipLimitError):
        _convert(zip_bytes, max_documents=1)


def test_traversing_member_rejected() -> None:
    zip_bytes = _make_zip([("R5-123456.docx", b"main"), ("../evil.docx", b"evil")])

    with pytest.raises(TDocZipValidationError):
        _convert(zip_bytes)


def test_main_inside_folder_keeps_full_path() -> None:
    zip_bytes = _make_zip([("R5-123456/R5-123456.docx", b"main"), ("annex/a.docx", b"annex")])

    result = _convert(zip_bytes)

    assert result.main_source_file == "R5-123456/R5-123456.docx"
    assert result.documents[0].source_file == "R5-123456/R5-123456.docx"


def test_is_empty_text_error_is_public() -> None:
    from doc3gpp.parsers.tdoc_zip_convert import is_empty_text_error

    assert callable(is_empty_text_error)


def test_is_empty_text_error_matches_empty_text_runtime_error() -> None:
    from doc3gpp.parsers.tdoc_zip_convert import is_empty_text_error

    assert is_empty_text_error(RuntimeError("python-docx returned empty result for x.docx"))


def test_is_empty_text_error_rejects_other_runtime_error() -> None:
    from doc3gpp.parsers.tdoc_zip_convert import is_empty_text_error

    assert not is_empty_text_error(RuntimeError("other failure"))


def test_is_empty_text_error_rejects_non_runtime_error() -> None:
    from doc3gpp.parsers.tdoc_zip_convert import is_empty_text_error

    assert not is_empty_text_error(ValueError("python-docx returned empty result for x.docx"))


def test_direct_extractor_does_not_import_private_empty_text_prefix() -> None:
    from doc3gpp.parsers import direct_extractor

    source = Path(direct_extractor.__file__).read_text(encoding="utf-8")

    assert "_EMPTY_TEXT_PREFIX" not in source


class _FakePool:
    """Records each ``convert_members`` call and returns the canned outcomes."""

    def __init__(self, *outcomes: object) -> None:
        self._outcomes = outcomes
        self.calls: list[tuple[str, list[tuple[str, bytes]]]] = []

    def convert_members(self, kind: str, members: Sequence[tuple[str, bytes]]) -> list[object]:
        self.calls.append((kind, list(members)))
        return list(self._outcomes)


def _convert_with_pool(zip_bytes: bytes, pool: _FakePool) -> TDocZipConversion:
    return convert_tdoc_zip(
        zip_bytes,
        tdoc_id="R5-123456",
        identity=IDENTITY,
        max_documents=100,
        max_uncompressed_bytes=0,
        docx_pool=pool,
    )


def test_docx_pool_annex_error_is_skipped_and_main_kept(
    caplog: pytest.LogCaptureFixture,
) -> None:
    zip_bytes = _make_zip([("R5-123456.docx", b"main body"), ("annex/a.docx", b"annex a body")])
    pool = _FakePool("main body", RuntimeError("cannot convert annex/a.docx"))

    with caplog.at_level(logging.WARNING):
        result = _convert_with_pool(zip_bytes, pool)

    assert [document.source_file for document in result.documents] == ["R5-123456.docx"]
    assert result.documents[0].markdown_text == "main body"
    assert result.skipped_sources == ("annex/a.docx",)
    assert any(
        record.levelno == logging.WARNING and "annex/a.docx" in record.getMessage()
        for record in caplog.records
    )


def test_docx_pool_main_error_raises() -> None:
    zip_bytes = _make_zip([("R5-123456.docx", b"main body"), ("annex/a.docx", b"annex a body")])
    pool = _FakePool(RuntimeError("cannot convert R5-123456.docx"), "annex a body")

    with pytest.raises(RuntimeError, match="cannot convert R5-123456.docx"):
        _convert_with_pool(zip_bytes, pool)


def test_docx_pool_empty_text_annex_is_kept_as_empty_markdown() -> None:
    zip_bytes = _make_zip(
        [
            ("R5-123456.docx", b"main body"),
            ("annex/a.docx", b"unused"),
            ("annex/b.docx", b"annex b body"),
        ]
    )
    pool = _FakePool(
        "main body",
        RuntimeError("python-docx returned empty result for annex/a.docx"),
        "annex b body",
    )

    result = _convert_with_pool(zip_bytes, pool)

    assert [document.source_file for document in result.documents] == [
        "R5-123456.docx",
        "annex/a.docx",
        "annex/b.docx",
    ]
    assert [document.markdown_text for document in result.documents] == [
        "main body",
        "",
        "annex b body",
    ]
    assert result.skipped_sources == ()


def test_docx_pool_empty_text_main_is_kept_as_empty_markdown() -> None:
    zip_bytes = _make_zip([("R5-123456.docx", b"unused"), ("annex/a.docx", b"annex a body")])
    pool = _FakePool(
        RuntimeError("python-docx returned empty result for R5-123456.docx"),
        "annex a body",
    )

    result = _convert_with_pool(zip_bytes, pool)

    assert result.documents[0].markdown_text == ""
    assert result.documents[0].content_hash == hashlib.sha256(b"").hexdigest()
    assert result.skipped_sources == ()


def test_docx_pool_receives_markdown_members_main_first() -> None:
    zip_bytes = _make_zip(
        [
            ("annex/b.docx", b"annex b body"),
            ("R5-123456.docx", b"main body"),
            ("annex/a.docx", b"annex a body"),
        ]
    )
    pool = _FakePool("main body", "annex a body", "annex b body")

    _convert_with_pool(zip_bytes, pool)

    assert pool.calls == [
        (
            "markdown",
            [
                ("R5-123456.docx", b"main body"),
                ("annex/a.docx", b"annex a body"),
                ("annex/b.docx", b"annex b body"),
            ],
        )
    ]


def test_docx_pool_is_not_used_with_injected_convert() -> None:
    zip_bytes = _make_zip([("R5-123456.docx", b"main body"), ("annex/a.docx", b"annex a body")])
    pool = _FakePool()

    result = convert_tdoc_zip(
        zip_bytes,
        tdoc_id="R5-123456",
        identity=IDENTITY,
        max_documents=100,
        max_uncompressed_bytes=0,
        convert=_decode,
        docx_pool=pool,
    )

    assert pool.calls == []
    assert [document.markdown_text for document in result.documents] == [
        "main body",
        "annex a body",
    ]


def test_docx_pool_error_on_annex_is_not_skipped() -> None:
    zip_bytes = _make_zip([("R5-123456.docx", b"main body"), ("annex/a.docx", b"annex a body")])
    pool = _FakePool(
        "main body",
        DocxConvertPoolError("DOCX conversion pool broke while converting annex/a.docx"),
    )

    with pytest.raises(DocxConvertPoolError, match="pool broke"):
        _convert_with_pool(zip_bytes, pool)


def test_docx_pool_error_on_main_raises() -> None:
    zip_bytes = _make_zip([("R5-123456.docx", b"main body"), ("annex/a.docx", b"annex a body")])
    pool = _FakePool(
        DocxConvertPoolError("DOCX conversion pool broke while converting R5-123456.docx"),
        "annex a body",
    )

    with pytest.raises(DocxConvertPoolError, match="pool broke"):
        _convert_with_pool(zip_bytes, pool)


def test_docx_pool_serves_direct_main_markdown_sniff() -> None:
    zip_bytes = _make_zip([("R5-123456.docx", b"main body"), ("annex/a.docx", b"annex a body")])
    pool = _FakePool("main body")

    markdown_text = convert_direct_main_markdown(
        zip_bytes,
        filename="R5-123456.zip",
        tdoc_id="R5-123456",
        max_documents=100,
        max_uncompressed_bytes=0,
        docx_pool=pool,
    )

    assert markdown_text == "main body"
    assert pool.calls == [("markdown", [("R5-123456.docx", b"main body")])]


@pytest.mark.skipif(
    not _docx_available(),
    reason="python-docx not installed; install with `pip install doc3gpp[extract]`",
)
def test_docx_pool_is_not_used_for_standalone_docx() -> None:
    pool = _FakePool()

    conversion = convert_direct_payload(
        _make_docx("standalone body"),
        filename="R5-123456.docx",
        tdoc_id="R5-123456",
        identity=IDENTITY,
        max_documents=100,
        max_uncompressed_bytes=0,
        docx_pool=pool,
    )

    assert pool.calls == []
    assert "standalone body" in conversion.documents[0].markdown_text

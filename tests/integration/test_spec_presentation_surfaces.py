from __future__ import annotations

import asyncio
import json
import shutil
from urllib.parse import parse_qs, quote, urlsplit

import pytest
from bs4 import BeautifulSoup
from fastapi.testclient import TestClient
from mcp.server.mcpserver.exceptions import ToolError
from typer.testing import CliRunner

from doc3gpp.cli import app
from doc3gpp.models.spec import Spec, SpecVersion
from doc3gpp.models.spec_doc import SpecDocPresentation
from doc3gpp.services.spec_doc_service import SpecDocService
from doc3gpp.settings.schema import (
    CacheSettings,
    MCPSettings,
    ServerSettings,
    Settings,
    SpecDocSettings,
)
from doc3gpp.storage.db.migrate import create_schema
from doc3gpp.storage.repositories.spec_doc_sql import SQLAlchemySpecDocRepository
from doc3gpp.storage.repositories.spec_sql import SQLAlchemySpecRepository
from doc3gpp.web.app import build_app, build_state
from doc3gpp.web.mcp_server import build_mcp_server

SPEC_ID = "38.331"
VERSION = "18.5.0"
SOURCE_FILE = 'slides/<script>&"日本語.pptx'
FIRST_SOURCE_FILE = "slides/first.pptx"
MARKDOWN = "# Presentation\n\n<script>alert(1)</script>"


@pytest.fixture
def presentation_corpus(sqlite_env, tmp_path, monkeypatch):
    create_schema("all")
    spec_repo = SQLAlchemySpecRepository()
    spec_repo.upsert(Spec(spec_id=SPEC_ID, type="TS", title="NR RRC"))
    spec_repo.upsert_versions(
        [
            SpecVersion(
                spec_id=SPEC_ID,
                version=VERSION,
                ftp_url="https://example.test/38.331-18.5.0.zip",
                release="Rel-18",
            )
        ]
    )

    doc_repo = SQLAlchemySpecDocRepository()
    doc_repo.record_download(
        SPEC_ID,
        VERSION,
        release="Rel-18",
        ftp_url="https://example.test/38.331-18.5.0.zip",
        docx_count=2,
    )
    doc_repo.replace_version_content(
        SPEC_ID,
        VERSION,
        release="Rel-18",
        drafts=[],
        presentations=[
            SpecDocPresentation(
                spec_id=SPEC_ID,
                version=VERSION,
                source_file=FIRST_SOURCE_FILE,
                file_order=0,
                markdown="# First presentation",
            ),
            SpecDocPresentation(
                spec_id=SPEC_ID,
                version=VERSION,
                source_file=SOURCE_FILE,
                file_order=1,
                markdown=MARKDOWN,
            ),
        ],
    )
    doc_repo.record_parsed(SPEC_ID, VERSION, chunk_count=0)

    cache_dir = tmp_path / "spec-doc-cache"
    zip_cache = cache_dir / "zips" / SPEC_ID / f"{VERSION}.zip"
    zip_cache.parent.mkdir(parents=True)
    zip_cache.write_bytes(b"stale ZIP cache")
    markdown_cache = cache_dir / "markdown" / SPEC_ID / VERSION
    markdown_cache.mkdir(parents=True)
    (markdown_cache / "1-slides.md").write_text("stale Markdown cache", encoding="utf-8")
    shutil.rmtree(cache_dir)
    assert not cache_dir.exists()

    def unexpected_call(*_args, **_kwargs):
        raise AssertionError("presentation reads must not download, extract, or read caches")

    monkeypatch.setattr(
        "doc3gpp.services.spec_doc_service.fetch_spec_doc_zip", unexpected_call
    )
    monkeypatch.setattr(SpecDocService, "_read_zip_cache", unexpected_call)
    monkeypatch.setattr(SpecDocService, "prepare_parse", unexpected_call)
    monkeypatch.setattr(SpecDocService, "parse", unexpected_call)

    settings = Settings(
        server=ServerSettings(enabled=True, port=8765),
        mcp=MCPSettings(enabled=True),
        cache=CacheSettings(dir=tmp_path / "tdoc-cache"),
        spec_doc=SpecDocSettings(cache_dir=cache_dir),
    )

    def build_cli_service() -> SpecDocService:
        return SpecDocService(
            spec_repo=spec_repo,
            doc_repo=doc_repo,
            settings=settings,
            cache_dir=cache_dir,
        )

    monkeypatch.setattr("doc3gpp.cli.build_spec_doc_service", build_cli_service)
    return settings


def _show_args(*, version: str = VERSION, source_file: str = SOURCE_FILE) -> list[str]:
    return [
        "spec",
        "doc",
        "show",
        "--spec",
        SPEC_ID,
        "--version",
        version,
        "--source-file",
        source_file,
    ]


def _call_mcp(server, name: str, arguments: dict):
    async def call():
        return await server.call_tool(name, arguments)

    return asyncio.run(call())


def test_spec_doc_show_cli_raw_json_and_output_file(presentation_corpus, tmp_path):
    runner = CliRunner()

    raw = runner.invoke(app, _show_args())
    assert raw.exit_code == 0, raw.output
    assert raw.stdout == MARKDOWN

    compact = runner.invoke(
        app, [*_show_args(), "--format", "json", "--compact"]
    )
    assert compact.exit_code == 0, compact.output
    expected = {
        "spec_id": SPEC_ID,
        "version": VERSION,
        "source_file": SOURCE_FILE,
        "markdown": MARKDOWN,
    }
    assert json.loads(compact.stdout) == expected

    output_file = tmp_path / "presentation.json"
    written = runner.invoke(
        app,
        [
            *_show_args(),
            "--format",
            "json",
            "--compact",
            "--output",
            str(output_file),
        ],
    )
    assert written.exit_code == 0, written.output
    assert output_file.read_text(encoding="utf-8") == compact.stdout


def test_spec_doc_presentation_http_html_json_and_version_links(presentation_corpus):
    with TestClient(build_app(presentation_corpus)) as client:
        version_page = client.get(
            f"/specs/{SPEC_ID}/docs", params={"version": VERSION}
        )
        assert version_page.status_code == 200, version_page.text
        soup = BeautifulSoup(version_page.text, "html.parser")
        links = soup.select("#spec-doc-presentations a.spec-doc-presentation")
        assert [link.get_text() for link in links] == [
            FIRST_SOURCE_FILE,
            SOURCE_FILE,
        ]

        href = links[1]["href"]
        assert quote(SOURCE_FILE, safe="") in href
        assert "<" not in href
        assert '"' not in href
        assert "日本語" not in href
        assert parse_qs(urlsplit(href).query) == {
            "version": [VERSION],
            "source_file": [SOURCE_FILE],
        }

        content_page = client.get(href)
        assert content_page.status_code == 200, content_page.text
        assert SOURCE_FILE in BeautifulSoup(content_page.text, "html.parser").get_text()
        assert "<script>alert(1)</script>" not in content_page.text
        assert "&lt;script&gt;" in content_page.text

        json_response = client.get(
            f"/specs/{SPEC_ID}/docs/content",
            params={
                "version": VERSION,
                "source_file": SOURCE_FILE,
                "format": "json",
            },
        )

    assert json_response.status_code == 200, json_response.text
    assert json.loads(json_response.content) == {
        "spec_id": SPEC_ID,
        "version": VERSION,
        "source_file": SOURCE_FILE,
        "markdown": MARKDOWN,
    }


def test_mcp_get_spec_doc_content_returns_json(presentation_corpus):
    state = build_state(presentation_corpus)
    server = build_mcp_server(state)

    try:
        result = _call_mcp(
            server,
            "get_spec_doc_content",
            {"spec_id": SPEC_ID, "version": VERSION, "source_file": SOURCE_FILE},
        )
    except Exception as exc:  # noqa: BLE001 - report an absent or broken tool as an assertion
        pytest.fail(f"get_spec_doc_content failed: {exc}")

    assert result.is_error is False
    assert json.loads(result.content[0].text) == {
        "spec_id": SPEC_ID,
        "version": VERSION,
        "source_file": SOURCE_FILE,
        "markdown": MARKDOWN,
    }


def test_spec_doc_presentation_json_is_byte_identical_across_surfaces(
    presentation_corpus,
):
    cli = CliRunner().invoke(
        app,
        [*_show_args(), "--format", "json", "--compact"],
    )
    assert cli.exit_code == 0, cli.output

    with TestClient(build_app(presentation_corpus)) as client:
        http = client.get(
            f"/specs/{SPEC_ID}/docs/content",
            params={
                "version": VERSION,
                "source_file": SOURCE_FILE,
                "format": "json",
            },
        )
    assert http.status_code == 200, http.text

    state = build_state(presentation_corpus)
    server = build_mcp_server(state)
    mcp = _call_mcp(
        server,
        "get_spec_doc_content",
        {"spec_id": SPEC_ID, "version": VERSION, "source_file": SOURCE_FILE},
    )
    assert mcp.is_error is False
    mcp_json = mcp.content[0].text

    assert list(json.loads(cli.stdout)) == [
        "spec_id",
        "version",
        "source_file",
        "markdown",
    ]
    assert "日本語".encode() in http.content
    assert cli.stdout.encode("utf-8") == http.content == mcp_json.encode("utf-8")


@pytest.mark.parametrize(
    ("version", "source_file"),
    [(VERSION, "missing.pptx"), ("19.0.0", SOURCE_FILE)],
)
def test_spec_doc_show_cli_reports_missing_presentation(
    presentation_corpus, version: str, source_file: str
):
    result = CliRunner().invoke(
        app,
        [*_show_args(version=version, source_file=source_file), "--format", "json"],
    )

    assert result.exit_code != 0
    assert "presentation" in result.output
    assert "not found" in result.output
    assert source_file in result.output
    assert f"{SPEC_ID}@{version}" in result.output


@pytest.mark.parametrize(
    ("version", "source_file"),
    [(VERSION, "missing.pptx"), ("19.0.0", SOURCE_FILE)],
)
def test_spec_doc_presentation_http_reports_missing_as_404(
    presentation_corpus, version: str, source_file: str
):
    with TestClient(build_app(presentation_corpus)) as client:
        response = client.get(
            f"/specs/{SPEC_ID}/docs/content",
            params={"version": version, "source_file": source_file, "format": "json"},
        )

    assert response.status_code == 404
    detail = response.json()["detail"]
    assert "presentation" in detail
    assert "not found" in detail
    assert source_file in detail
    assert f"{SPEC_ID}@{version}" in detail


@pytest.mark.parametrize(
    ("version", "source_file"),
    [(VERSION, "missing.pptx"), ("19.0.0", SOURCE_FILE)],
)
def test_spec_doc_presentation_mcp_reports_missing_as_error(
    presentation_corpus, version: str, source_file: str
):
    server = build_mcp_server(build_state(presentation_corpus))
    with pytest.raises(ToolError) as exc_info:
        _call_mcp(
            server,
            "get_spec_doc_content",
            {"spec_id": SPEC_ID, "version": version, "source_file": source_file},
        )

    error_text = str(exc_info.value)
    assert "presentation" in error_text
    assert "not found" in error_text
    assert source_file in error_text
    assert f"{SPEC_ID}@{version}" in error_text


def test_parsed_mixed_zip_details_survive_cache_eviction_on_all_surfaces(
    sqlite_env, tmp_path, monkeypatch,
):
    import asyncio
    import io
    import json
    import shutil
    import zipfile
    from urllib.parse import parse_qs, quote, urlsplit

    from docx import Document
    from fastapi.testclient import TestClient
    from pptx import Presentation
    from pptx.util import Inches
    from typer.testing import CliRunner

    from doc3gpp.cli import app
    from doc3gpp.models.spec import Spec, SpecVersion
    from doc3gpp.parsers.presentation_converter import convert_presentation_to_markdown
    from doc3gpp.services.spec_doc_service import SpecDocService
    from doc3gpp.settings.schema import (
        CacheSettings,
        MCPSettings,
        ServerSettings,
        Settings,
        SpecDocSettings,
    )
    from doc3gpp.storage.db.migrate import create_schema
    from doc3gpp.storage.repositories.spec_doc_sql import SQLAlchemySpecDocRepository
    from doc3gpp.storage.repositories.spec_sql import SQLAlchemySpecRepository
    from doc3gpp.web.app import build_app, build_state
    from doc3gpp.web.mcp_server import build_mcp_server

    create_schema("all")
    spec_repo = SQLAlchemySpecRepository()
    spec_repo.upsert(Spec(spec_id=SPEC_ID, type="TS", title="NR RRC"))
    ftp_url = "https://example.test/38.331-18.5.0.zip"
    spec_repo.upsert_versions(
        [SpecVersion(SPEC_ID, VERSION, ftp_url, "Rel-18")]
    )

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, text in enumerate(
        ("Presentation", "<script>alert(1)</script>", "slide acceptance signal")
    ):
        shape = slide.shapes.add_textbox(
            Inches(1), Inches(index + 1), Inches(8), Inches(0.5)
        )
        shape.text_frame.text = text
    slide.notes_slide.notes_text_frame.text = "speaker note acceptance probe"
    pptx_buffer = io.BytesIO()
    presentation.save(pptx_buffer)
    pptx_bytes = pptx_buffer.getvalue()

    word = Document()
    word.add_heading("2 Scope", level=1)
    word.add_paragraph("Word section content.")
    word_buffer = io.BytesIO()
    word.save(word_buffer)

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as archive:
        archive.writestr(SOURCE_FILE, pptx_bytes)
        archive.writestr("38331-j00.docx", word_buffer.getvalue())
    zip_bytes = zip_buffer.getvalue()
    expected_markdown = convert_presentation_to_markdown(pptx_bytes, SOURCE_FILE)
    assert expected_markdown == (
        "## Slide 1\n\n"
        "Presentation\n\n"
        "<script>alert(1)</script>\n\n"
        "slide acceptance signal\n\n"
        "**Speaker notes:**\n\n"
        "speaker note acceptance probe"
    )

    doc_repo = SQLAlchemySpecDocRepository()
    cache_dir = tmp_path / "parsed-spec-doc-cache"
    settings = Settings(
        server=ServerSettings(enabled=True, port=8765),
        mcp=MCPSettings(enabled=True),
        cache=CacheSettings(dir=tmp_path / "tdoc-cache"),
        spec_doc=SpecDocSettings(cache_dir=cache_dir),
    )
    service = SpecDocService(
        spec_repo=spec_repo,
        doc_repo=doc_repo,
        settings=settings,
        fetcher=lambda _url: zip_bytes,
        cache_dir=cache_dir,
    )
    source = service.parse(SPEC_ID)
    stored = service.get_presentation(SPEC_ID, VERSION, SOURCE_FILE)
    assert source.docx_count == 2
    assert stored.markdown == expected_markdown

    shutil.rmtree(cache_dir)
    assert not cache_dir.exists()

    def unexpected_call(*_args, **_kwargs):
        raise AssertionError("presentation detail reads must not fetch, parse, or read caches")

    monkeypatch.setattr(
        "doc3gpp.services.spec_doc_service.fetch_spec_doc_zip", unexpected_call
    )
    monkeypatch.setattr(SpecDocService, "_read_zip_cache", unexpected_call)
    monkeypatch.setattr(SpecDocService, "prepare_parse", unexpected_call)
    monkeypatch.setattr(SpecDocService, "parse", unexpected_call)
    monkeypatch.setattr("doc3gpp.cli.build_spec_doc_service", lambda: service)

    cli_raw = CliRunner().invoke(app, _show_args())
    assert cli_raw.exit_code == 0, cli_raw.output
    assert cli_raw.stdout == expected_markdown
    cli_json = CliRunner().invoke(
        app, [*_show_args(), "--format", "json", "--compact"]
    )
    assert cli_json.exit_code == 0, cli_json.output

    with TestClient(build_app(settings)) as client:
        version_page = client.get(
            f"/specs/{SPEC_ID}/docs", params={"version": VERSION}
        )
        assert version_page.status_code == 200, version_page.text
        links = BeautifulSoup(version_page.text, "html.parser").select(
            "#spec-doc-presentations a.spec-doc-presentation"
        )
        assert [link.get_text() for link in links] == [SOURCE_FILE]
        href = links[0]["href"]
        assert quote(SOURCE_FILE, safe="") in href
        assert parse_qs(urlsplit(href).query) == {
            "version": [VERSION],
            "source_file": [SOURCE_FILE],
        }
        html = client.get(href)
        assert html.status_code == 200, html.text
        assert "<script>alert(1)</script>" not in html.text
        assert "&lt;script&gt;" in html.text
        http = client.get(
            f"/specs/{SPEC_ID}/docs/content",
            params={
                "version": VERSION,
                "source_file": SOURCE_FILE,
                "format": "json",
            },
        )
    assert http.status_code == 200, http.text
    assert http.json()["markdown"] == expected_markdown

    state = build_state(settings)
    try:
        server = build_mcp_server(state)
        mcp = asyncio.run(
            server.call_tool(
                "get_spec_doc_content",
                {"spec_id": SPEC_ID, "version": VERSION, "source_file": SOURCE_FILE},
            )
        )
        assert mcp.is_error is False
        mcp_json = mcp.content[0].text
    finally:
        if state.close_docx_pool is not None:
            state.close_docx_pool()
        if state.close_embedder is not None:
            state.close_embedder()
        state.engine.dispose()
        state.testcase_engine.dispose()
        if state.specdata_engine is not None:
            state.specdata_engine.dispose()

    assert json.loads(cli_json.stdout)["markdown"] == expected_markdown
    assert cli_json.stdout.encode("utf-8") == http.content == mcp_json.encode("utf-8")
    assert not cache_dir.exists()
